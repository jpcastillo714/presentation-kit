"""kit-presentaciones: un solo script para todo lo mecánico / one script for all the mechanics.

ES. El diseño no vive aquí: vive en el perfil (colores, tipografías, logo, plantilla) y en el guion (qué dice cada
diapositiva). Este script traduce el guion a HTML y a .pptx, empaqueta, exporta y revisa.
EN. Design does not live here: it lives in the profile (colors, fonts, logo, template) and in the script (what each
slide says). This tool turns the script into HTML and .pptx, packages, exports and reviews.

Comandos (alias en inglés entre paréntesis) / Commands (English aliases in parentheses):

  python motor/kit.py nueva (new) <carpeta> [--perfil|--profile sobrio] [--titulo|--title "..."]
      Crea una presentación con el motor, el perfil y un guion de partida. / Scaffolds a new deck.
  python motor/kit.py construir (build) <carpeta> [--html] [--pptx] [--guion|--script guion.en.md]
      Del guion a deck.html y <nombre>.pptx (guion.en.md -> deck.en.html, <nombre>.en.pptx). / Script -> HTML + .pptx.
  python motor/kit.py empaquetar (package) <deck.html> [salida.html]
      Un solo archivo HTML que funciona sin internet. / Single self-contained HTML file.
  python motor/kit.py pdf <archivo.html> [salida.pdf] [--sobrescribir|--overwrite]
      PDF con Chrome o Edge, una diapositiva por página. / PDF via Chrome or Edge, one slide per page.
  python motor/kit.py perfil (profile) --desde|--from ref.pptx --nombre|--name <n>  |  --regenerar|--regenerate <dir>
      Perfil a partir de una presentación de referencia. / Profile from a reference deck.
  python motor/kit.py revisar (review) <deck.html>
      Revisión mecánica contra los principios. / Mechanical review against the principles.
  python motor/kit.py capturas (shots) <deck.html> <carpeta_png> [--solo|--only 3,5]
      Una imagen por diapositiva para revisarla. / One PNG per slide, to review by eye.
  python motor/kit.py estaticos (statics) <carpeta> [--guion|--script guion.en.md]
      Estado final de cada animación como imagen, para el .pptx. / Final frame of each animation, for .pptx.
  python motor/kit.py renderizar (render) <archivo.pptx> <carpeta_png>
      Cada diapositiva del .pptx a PNG con PowerPoint (solo Windows). / .pptx -> PNG via PowerPoint (Windows only).
"""
from __future__ import annotations

import base64
import html
import mimetypes
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import yaml

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

KIT = Path(__file__).resolve().parent.parent          # raíz del kit / kit root
MOTOR = KIT / "motor"
PERFILES = KIT / "perfiles"
NAVEGADORES = [r"C:\Program Files\Google\Chrome\Application\chrome.exe",
               r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
               r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
               r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
               "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
               "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
               "/Applications/Chromium.app/Contents/MacOS/Chromium"]


def navegador() -> str:
    """Chrome, Edge o Chromium sin ventana (PDF, capturas). Variable KIT_NAVEGADOR para forzar uno.
    Chrome, Edge or Chromium in headless mode. Set KIT_NAVEGADOR to force a specific binary."""
    if os.environ.get("KIT_NAVEGADOR"):
        return os.environ["KIT_NAVEGADOR"]
    for ruta in NAVEGADORES:
        if Path(ruta).exists():
            return ruta
    for nombre in ("google-chrome", "google-chrome-stable", "chromium", "chromium-browser", "microsoft-edge",
                   "msedge", "chrome"):
        if shutil.which(nombre):
            return shutil.which(nombre)
    sys.exit("[ERROR] No encontré Chrome, Edge ni Chromium (define KIT_NAVEGADOR). / No Chrome, Edge or Chromium found.")


def _sufijo(nombre_guion: str) -> str:
    """guion.md -> '' ; guion.en.md -> '.en'."""
    m = re.match(r"guion(.*)\.md$", nombre_guion)
    return m.group(1) if m else ""


# =====================================================================
# Perfiles
# =====================================================================

def leer_perfil(carpeta: Path) -> dict:
    datos = yaml.safe_load((carpeta / "perfil.yml").read_text(encoding="utf-8"))
    datos["_carpeta"] = carpeta
    return datos


def css_de_perfil(p: dict) -> str:
    c, f, t = p["colores"], p["fuentes"], p.get("tamanos", {})
    caras = []
    carpeta_fuentes = p["_carpeta"] / "fuentes"
    if carpeta_fuentes.is_dir():
        for arch in sorted(carpeta_fuentes.iterdir()):
            if arch.suffix.lower() not in (".woff2", ".woff", ".ttf", ".otf"):
                continue
            # Convención de nombre: Familia-Peso[-Italic].ext, p. ej. SourceSerif4-600.woff2
            m = re.match(r"(.+?)-(\d{3})(-italic)?$", arch.stem, re.I)
            familia, peso, it = (m.group(1), m.group(2), m.group(3)) if m else (arch.stem, "400", None)
            caras.append(f"@font-face {{ font-family: '{familia}'; src: url('fuentes/{arch.name}'); "
                         f"font-weight: {peso}; font-style: {'italic' if it else 'normal'}; }}")
    v = {
        "--fondo": c["fondo"], "--tinta": c["tinta"], "--tinta-suave": c["tinta_suave"], "--acento": c["acento"],
        "--acento-2": c["acento_2"], "--linea": c["linea"], "--marco": c.get("marco", "#111111"),
        "--fuente-titulo": f["titulo"], "--fuente-texto": f["texto"], "--fuente-mono": f.get("mono", "Consolas, monospace"),
        "--titulo": f"{t.get('titulo', 60)}px", "--texto": f"{t.get('texto', 34)}px",
        "--peso-titulo": str(t.get("peso_titulo", 600)),
    }
    for i, s in enumerate(c.get("serie", []), 1):
        v[f"--serie-{i}"] = s
    cuerpo = "\n".join(f"  {k}: {val};" for k, val in v.items())
    extra = (p.get("css_extra") or "").strip()
    if p.get("fondo_imagen"):
        extra = f".slide {{ background-image: url('{p['fondo_imagen']}'); background-size: cover; }}\n" + extra
    return (f"/* Perfil «{p['nombre']}» — generado desde perfil.yml por kit.py. Edita el .yml, no este archivo. */\n"
            + "\n".join(caras) + ("\n" if caras else "") + f":root {{\n{cuerpo}\n}}\n" + (extra + "\n" if extra else ""))


def mplstyle_de_perfil(p: dict) -> str:
    c, f = p["colores"], p["fuentes"]
    hexes = ", ".join(f"'{s.lstrip('#')}'" for s in c.get("serie", [c["acento"], c["acento_2"]]))
    familias = ", ".join(f.get("figuras", ["DejaVu Sans"]))
    q = lambda s: s.lstrip("#")
    # Tamaños pensados para una figura que ocupa el cuerpo de una diapositiva de 1920×1080 (principio 14).
    return f"""# Estilo de figuras del perfil «{p['nombre']}» — generado desde perfil.yml por kit.py.
figure.figsize : 12, 5.4
figure.dpi : 100
savefig.dpi : 160
savefig.bbox : tight
savefig.facecolor : {q(c['fondo'])}
figure.facecolor : {q(c['fondo'])}
axes.facecolor : {q(c['fondo'])}
axes.edgecolor : {q(c['tinta_suave'])}
axes.labelcolor : {q(c['tinta'])}
axes.prop_cycle : cycler('color', [{hexes}])
axes.spines.top : False
axes.spines.right : False
axes.grid : False
axes.titlesize : 18
axes.labelsize : 17
axes.titlelocation : left
xtick.color : {q(c['tinta_suave'])}
ytick.color : {q(c['tinta_suave'])}
xtick.labelsize : 15
ytick.labelsize : 15
text.color : {q(c['tinta'])}
font.family : sans-serif
font.sans-serif : {familias}, DejaVu Sans
font.size : 16
legend.frameon : False
legend.fontsize : 15
lines.linewidth : 2.4
lines.markersize : 7
image.cmap : cividis
svg.fonttype : none
"""


def regenerar_perfil(carpeta: Path) -> None:
    p = leer_perfil(carpeta)
    (carpeta / "perfil.css").write_text(css_de_perfil(p), encoding="utf-8")
    (carpeta / "figuras.mplstyle").write_text(mplstyle_de_perfil(p), encoding="utf-8")
    print(f"[OK] perfil «{p['nombre']}»: perfil.css y figuras.mplstyle regenerados en {carpeta}")


NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "rel": "http://schemas.openxmlformats.org/package/2006/relationships"}


def _color_tema(el) -> str | None:
    if el is None:
        return None
    s = el.find("a:srgbClr", NS)
    if s is not None:
        return "#" + s.get("val").lower()
    s = el.find("a:sysClr", NS)
    if s is not None and s.get("lastClr"):
        return "#" + s.get("lastClr").lower()
    return None


def _luminancia(hexa: str) -> float:
    r, g, b = (int(hexa[i:i + 2], 16) / 255 for i in (1, 3, 5))
    lin = lambda x: x / 12.92 if x <= 0.03928 else ((x + 0.055) / 1.055) ** 2.4
    return 0.2126 * lin(r) + 0.7152 * lin(g) + 0.0722 * lin(b)


def contraste(a: str, b: str) -> float:
    la, lb = sorted((_luminancia(a), _luminancia(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


def _mezcla(a: str, b: str, t: float) -> str:
    ca = [int(a[i:i + 2], 16) for i in (1, 3, 5)]
    cb = [int(b[i:i + 2], 16) for i in (1, 3, 5)]
    return "#" + "".join(f"{round(x + (y - x) * t):02x}" for x, y in zip(ca, cb))


def perfil_desde_pptx(ref: Path, nombre: str, destino: Path) -> Path:
    """Lee el tema de una presentación de referencia y arma un perfil. Lo que no se puede leer, se avisa."""
    carpeta = destino / nombre
    if carpeta.exists():
        sys.exit(f"[ERROR] Ya existe el perfil {carpeta}. Elige otro nombre o bórralo tú.")
    avisos = []
    with zipfile.ZipFile(ref) as z:
        nombres = z.namelist()
        temas = sorted(n for n in nombres if re.match(r"ppt/theme/theme\d+\.xml$", n))
        if not temas:
            sys.exit("[ERROR] La referencia no tiene tema: ¿es un .pptx?")
        tema = ET.fromstring(z.read(temas[0]))
        esquema = tema.find(".//a:clrScheme", NS)
        col = {hijo.tag.split("}")[1]: _color_tema(hijo) for hijo in esquema} if esquema is not None else {}
        mayor = tema.find(".//a:majorFont/a:latin", NS)
        menor = tema.find(".//a:minorFont/a:latin", NS)
        f_tit = mayor.get("typeface") if mayor is not None else "Georgia"
        f_txt = menor.get("typeface") if menor is not None else "Segoe UI"
        # Imágenes que se repiten en muchas diapositivas o están en el patrón: probablemente logos.
        uso: dict[str, int] = {}
        rels_slides = [n for n in nombres if re.match(r"ppt/slides/_rels/slide\d+\.xml\.rels$", n)]
        rels_patron = [n for n in nombres if re.match(r"ppt/(slideMasters|slideLayouts)/_rels/.+\.rels$", n)]
        for n in rels_slides + rels_patron:
            for r in ET.fromstring(z.read(n)).findall("rel:Relationship", NS):
                t = r.get("Target", "")
                if "/media/" in t or t.startswith("../media/"):
                    clave = "ppt/media/" + t.split("media/")[1]
                    uso[clave] = uso.get(clave, 0) + (100 if n in rels_patron else 1)
        n_slides = max(1, len(rels_slides))
        # Colores que de verdad usan las diapositivas (rellenos y texto): el tema puede ser el de Office por
        # defecto aunque la presentación use otra paleta (pasa con presentaciones generadas por código).
        usados: dict[str, int] = {}
        for n in nombres:
            if re.match(r"ppt/(slides/slide|slideMasters/slideMaster|slideLayouts/slideLayout)\d+\.xml$", n):
                for val in re.findall(rb'<a:srgbClr val="([0-9A-Fa-f]{6})"', z.read(n)):
                    k = "#" + val.decode().lower()
                    usados[k] = usados.get(k, 0) + 1
        colores_usados = [k for k, _ in sorted(usados.items(), key=lambda x: -x[1])[:10]]
        logos = [m for m, k in sorted(uso.items(), key=lambda x: -x[1]) if k >= 100 or k >= 0.5 * n_slides]
        carpeta.mkdir(parents=True)
        copiados, fondos = [], []
        for m in logos[:6]:
            if m in nombres and Path(m).suffix.lower() in (".png", ".jpg", ".jpeg", ".svg", ".gif", ".emf", ".wmf"):
                datos_img = z.read(m)
                es_fondo = False
                try:   # una imagen grande con la forma de la diapositiva es un fondo, no un logo
                    from io import BytesIO
                    from PIL import Image
                    with Image.open(BytesIO(datos_img)) as im:
                        w_img, h_img = im.size
                    es_fondo = w_img >= 900 and 1.2 <= w_img / h_img <= 1.9
                except Exception:
                    pass
                sub = "fondos" if es_fondo else "logos"
                (carpeta / sub).mkdir(exist_ok=True)
                (carpeta / sub / Path(m).name).write_bytes(datos_img)
                (fondos if es_fondo else copiados).append(Path(m).name)
    fondo = col.get("lt1") or "#ffffff"
    tinta = col.get("dk1") or "#1a1a1a"
    fondo_imagen = None
    if fondos:
        # El fondo es una imagen puesta en las diapositivas (no en el patrón): se usa como fondo del perfil, y el
        # color de fondo y la tinta se ajustan a su tono para que el texto siga contrastando.
        from PIL import Image
        with Image.open(carpeta / "fondos" / fondos[0]) as im:
            r_, g_, b_ = im.convert("RGB").resize((1, 1), Image.Resampling.BOX).getpixel((0, 0))
        fondo = f"#{r_:02x}{g_:02x}{b_:02x}"
        fondo_imagen = f"fondos/{fondos[0]}"
        if _luminancia(fondo) < 0.2:
            claros = [c_ for c_ in colores_usados if _luminancia(c_) > 0.7]
            tinta = claros[0] if claros else "#f2f2f2"
            avisos.append(f"El fondo es una imagen oscura ({fondo}): la tinta pasó a {tinta}. Las figuras de "
                          "matplotlib con este perfil saldrán con fondo oscuro.")
        col["lt2"] = None
    tinta_suave = col.get("dk2") or _mezcla(tinta, fondo, 0.35)
    if contraste(tinta_suave, fondo) < 4.5:
        tinta_suave = _mezcla(tinta, fondo, 0.35)
        avisos.append("dk2 del tema no contrasta lo suficiente con el fondo; tinta_suave se calculó mezclando.")
    acentos = [col.get(f"accent{i}") for i in range(1, 7) if col.get(f"accent{i}")]
    if not acentos:
        acentos = ["#1f5f8b", "#b4532a"]
        avisos.append("El tema no trae acentos; se usaron los del perfil sobrio.")
    for a in acentos[:2]:
        if contraste(a, fondo) < 3:
            avisos.append(f"El acento {a} contrasta poco con el fondo {fondo} (< 3:1): úsalo solo en áreas grandes.")
    linea = col.get("lt2") or _mezcla(fondo, tinta, 0.15)
    if col.get("accent1") == "#4472c4" and col.get("accent2") == "#ed7d31":
        avisos.append("El tema es el de Office por defecto: lo más probable es que la paleta real esté en "
                      "colores_usados. Asigna fondo, tinta y acentos desde ahí y corre «kit.py perfil --regenerar».")
    shutil.copy2(ref, carpeta / "plantilla.pptx")
    datos = {
        "nombre": nombre,
        "descripcion": f"Extraído de {ref.name}. Revisa colores y logo antes de usarlo.",
        "origen": str(ref),
        "colores": {"fondo": fondo, "tinta": tinta, "tinta_suave": tinta_suave, "acento": acentos[0],
                    "acento_2": acentos[1] if len(acentos) > 1 else _mezcla(acentos[0], tinta, 0.4),
                    "linea": linea, "marco": "#111111", "serie": acentos[:5]},
        "fuentes": {"titulo": f"'{f_tit}', Georgia, serif", "texto": f"'{f_txt}', 'Segoe UI', system-ui, sans-serif",
                    "mono": "Consolas, monospace", "pptx_titulo": f_tit, "pptx_texto": f_txt,
                    "figuras": [f_txt, "Segoe UI"]},
        "tamanos": {"titulo": 60, "texto": 34, "peso_titulo": 600},
        "logo": f"logos/{copiados[0]}" if copiados else None,
        "fondo_imagen": fondo_imagen,
        "plantilla_pptx": "plantilla.pptx",
        "layout_pptx": None,
        "colores_usados": colores_usados,
        "avisos": avisos,
    }
    (carpeta / "perfil.yml").write_text(yaml.safe_dump(datos, allow_unicode=True, sort_keys=False), encoding="utf-8")
    regenerar_perfil(carpeta)
    print(f"[OK] perfil «{nombre}» en {carpeta}")
    print(f"     colores del tema: {', '.join(f'{k}={v}' for k, v in col.items() if v)}")
    print(f"     fuentes del tema: títulos «{f_tit}», texto «{f_txt}»")
    print(f"     logos candidatos: {', '.join(copiados) or 'ninguno'}")
    if fondos:
        print(f"     imagen de fondo de las diapositivas (fondo_imagen del perfil): {', '.join(fondos)}")
    print(f"     colores más usados en las diapositivas: {', '.join(colores_usados) or 'ninguno explícito'}")
    for a in avisos:
        print(f"[AVISO] {a}")
    print("     Mira el perfil antes de usarlo: el tema dice qué colores ofrece, no cuáles usa la presentación.")
    return carpeta


# =====================================================================
# Guion
# =====================================================================

CLAVES_EN = {"type": "tipo", "figure": "figura", "source": "fuente", "number": "cifra", "static": "estatico",
             "subtitle": "subtitulo", "steps": "pasos", "appendix": "anexo", "why": "porque", "class": "clase",
             "kicker": "antetitulo"}
TIPOS_EN = {"cover": "portada", "section": "seccion", "figure": "figura", "two": "dos", "number": "dato", "text": "texto"}
META_EN = {"title": "titulo", "author": "autor", "date": "fecha", "place": "lugar", "profile": "perfil",
           "transition": "transicion", "motion": "movimiento", "lang": "idioma", "language": "idioma"}
VALORES_EN = {"fade": "fundido", "respect": "respetar", "full": "completa"}


def leer_guion(ruta: Path) -> tuple[dict, list[dict]]:
    """guion.md: front matter YAML + una diapositiva por «## Título». Ver ejemplos/guion_ejemplo.md."""
    texto = ruta.read_text(encoding="utf-8")
    meta = {}
    m = re.match(r"^---\n(.*?)\n---\n", texto, re.S)
    if m:
        meta = yaml.safe_load(m.group(1)) or {}
        meta = {META_EN.get(k, k): (VALORES_EN.get(v, v) if isinstance(v, str) else v) for k, v in meta.items()}
        texto = texto[m.end():]
    slides, sl = [], None
    for linea in texto.splitlines():
        if linea.startswith("## "):
            sl = {"titulo": linea[3:].strip(), "tipo": None, "figuras": [], "vinetas": [], "parrafos": [],
                  "notas": [], "fuente": None, "anim": None, "estatico": None, "cifra": None, "anexo": False,
                  "pasos": False, "subtitulo": None, "porque": None, "clase": None, "antetitulo": None}
            slides.append(sl)
            continue
        if sl is None or not linea.strip() or linea.startswith("<!--"):
            continue
        s = linea.strip()
        if s.startswith("@"):
            clave, _, valor = s[1:].partition(" ")
            valor = valor.strip()
            clave = CLAVES_EN.get(clave, clave)
            if clave == "tipo":
                valor = TIPOS_EN.get(valor, valor)
            if clave == "clase":
                valor = " ".join(VALORES_EN.get(v, v) for v in valor.split())
            if clave == "figura":
                sl["figuras"].append(valor)
            elif clave in ("anexo", "pasos"):
                sl[clave] = int(valor) if valor.isdigit() else True   # «@pasos 3»: niveles de una animación
            elif clave in sl:
                sl[clave] = valor
            else:
                print(f"[AVISO] clave desconocida @{clave} en «{sl['titulo']}»")
        elif s.startswith(">"):
            sl["notas"].append(s.lstrip("> ").strip())
        elif s.startswith("- "):
            sl["vinetas"].append(s[2:].strip())
        else:
            sl["parrafos"].append(s)
    for sl in slides:
        if not sl["tipo"]:
            sl["tipo"] = ("dato" if sl["cifra"] else "dos" if len(sl["figuras"]) > 1
                          else "figura" if sl["figuras"] or sl["anim"] else "texto")
    return meta, slides


def _md(s: str) -> str:
    """Markdown mínimo dentro de una línea: **negrita**, *énfasis*, `código`, ^superíndice^."""
    s = html.escape(s, quote=False)
    s = re.sub(r"\^(.+?)\^", r"<sup>\1</sup>", s)
    s = re.sub(r"\*\*(.+?)\*\*", r'<b class="acento">\1</b>', s)
    s = re.sub(r"\*(.+?)\*", r"<em>\1</em>", s)
    return re.sub(r"`(.+?)`", r"<code>\1</code>", s)


_SUP = str.maketrans("0123456789−-+", "⁰¹²³⁴⁵⁶⁷⁸⁹⁻⁻⁺")


def _plano(s: str) -> str:
    """Texto para el .pptx: sin marcas de markdown; ^−10^ pasa a caracteres en superíndice."""
    s = re.sub(r"\^(.+?)\^", lambda m: m.group(1).translate(_SUP), s)
    return re.sub(r"\*\*|\*|`", "", s)


def html_de_guion(meta: dict, slides: list[dict], tiene_logo: bool, carpeta: Path | None = None) -> str:
    """deck.html se regenera desde el guion. Lo propio de cada presentación vive aparte y se enlaza si existe:
    estilos.css (ajustes de diseño) y animaciones.js (window.animaciones.<nombre>)."""
    partes = []
    for i, sl in enumerate(slides, 1):
        t, clases = sl["tipo"], ["slide", sl["tipo"]] + (sl["clase"].split() if sl["clase"] else [])
        if sl["anexo"]:
            clases.append("anexo")
        attrs = f' data-anim="{sl["anim"]}"' if sl["anim"] else ""
        if type(sl["pasos"]) is int:      # bool no cuenta: True y False también son int en Python
            attrs += f' data-pasos="{sl["pasos"]}"'
        paso = (lambda k: f' data-paso="{k}"') if sl["pasos"] is True else (lambda k: "")
        h = [f'<section class="{" ".join(clases)}"{attrs}>']
        if sl["anexo"]:
            h.append(f'  <div class="etiqueta-anexo">{"Appendix" if meta.get("idioma") == "en" else "Anexo"}</div>')
        if t == "portada":
            if sl["anim"]:
                h.append(f'  <div class="cuerpo"><figure><div class="lienzo" id="lienzo-{i}"></div></figure></div>')
            h.append(f"  <h1>{_md(sl['titulo'])}</h1>")
            if sl["subtitulo"]:
                h.append(f'  <p class="subtitulo">{_md(sl["subtitulo"])}</p>')
            aut = " · ".join(str(x) for x in (meta.get("autor"), meta.get("lugar"), meta.get("fecha")) if x)
            h.append(f'  <p class="autoria">{html.escape(aut)}</p>')
        elif t == "seccion":
            if sl["subtitulo"]:
                h.append(f'  <div class="numero-seccion">{_md(sl["subtitulo"])}</div>')
            h.append(f"  <h1>{_md(sl['titulo'])}</h1>")
        else:
            if sl["antetitulo"]:
                h.append(f'  <p class="antetitulo">{_md(sl["antetitulo"])}</p>')
            h.append(f"  <h1>{_md(sl['titulo'])}</h1>")
            texto = [f"<p{paso(k)}>{_md(p)}</p>" for k, p in enumerate(sl["parrafos"], 1)]
            if sl["vinetas"]:
                k0 = len(texto)
                texto.append("<ul>" + "".join(f"<li{paso(k0 + k)}>{_md(v)}</li>" for k, v in enumerate(sl["vinetas"], 1)) + "</ul>")
            figs = [f'<figure><img src="{html.escape(f)}" alt=""></figure>' for f in sl["figuras"]]
            if sl["anim"]:
                # En pantalla, el lienzo animado; en el PDF, la imagen estática (@estatico) si la hay.
                estatico = (f'<img class="solo-impreso" src="{html.escape(sl["estatico"])}" alt="">'
                            if sl["estatico"] else "")
                figs = [f'<figure><div class="lienzo{" solo-pantalla" if estatico else ""}" id="lienzo-{i}"><!-- animación «{sl["anim"]}»: '
                        f'dibújala en el <script> de animaciones --></div>{estatico}</figure>']
            if t == "dato":
                bloque = [f'<div class="lado"><div class="cifra">{_md(sl["cifra"])}</div></div>',
                          f'<div>{"".join(texto)}</div>']
                h.append('  <div class="cuerpo">' + "".join(bloque) + "</div>")
            elif t in ("figura", "dos") and texto and "completa" not in clases:
                h.append('  <div class="cuerpo">' + "".join(figs) + f'<div class="lado">{"".join(texto)}</div></div>')
            elif t in ("figura", "dos"):
                h.append(f'  <div class="cuerpo{" dos" if t == "dos" else ""}">' + "".join(figs) + "</div>")
            else:
                h.append('  <div class="cuerpo columna arriba">' + "".join(texto) + "</div>")
        if sl["fuente"]:
            h.append(f'  <p class="fuente">{_md(sl["fuente"])}</p>')
        if sl["notas"]:
            h.append('  <aside class="notas">' + " ".join(_md(n) for n in sl["notas"]) + "</aside>")
        if sl["porque"]:
            h.append(f'  <aside class="porque">{_md(sl["porque"])}</aside>')
        h.append("</section>")
        partes.append("\n".join(h))
    logo = ' data-logo="perfil/' + "{logo}" + '"' if tiene_logo else ""
    estilos = '\n<link rel="stylesheet" href="estilos.css">' if carpeta and (carpeta / "estilos.css").exists() else ""
    anim = ('\n<script src="motor/anim.js"></script>\n<script src="animaciones.js"></script>'
            if carpeta and (carpeta / "animaciones.js").exists() else "")
    if meta.get("transicion") == "fundido":
        logo += ' class="fundido"'
    # Las animaciones las ve la sala, no el computador de quien presenta: por defecto se anima aunque el sistema
    # pida reducir movimiento. «movimiento: respetar» en el guion devuelve la decisión al sistema.
    logo += f' data-movimiento="{"respetar" if meta.get("movimiento") == "respetar" else "siempre"}"'
    return f"""<!doctype html>
<!-- Generado desde guion por motor/kit.py. No lo edites a mano: cambia el guion y vuelve a construir.
     Ajustes de diseño en estilos.css; animaciones en animaciones.js (se enlazan solos si existen).
     Generated from the script by motor/kit.py. Do not edit by hand: edit the script and rebuild. -->
<html lang="{meta.get('idioma', 'es')}">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html.escape(str(meta.get('titulo', 'Presentación')))}</title>
<link rel="stylesheet" href="motor/deck.css">
<link rel="stylesheet" href="perfil/perfil.css">{estilos}
</head>
<body{logo}>
<div class="escenario">

{chr(10).join(partes)}

</div>
<script>window.animaciones = window.animaciones || {{}};</script>{anim}
<script src="motor/deck.js"></script>
</body>
</html>
"""


# =====================================================================
# .pptx
# =====================================================================

def _borrar_diapositivas(prs) -> None:
    lista = prs.slides._sldIdLst
    for sid in list(lista):
        prs.part.drop_rel(sid.rId)
        lista.remove(sid)


def _layout(prs, preferido: str | None):
    """El diseño de la plantilla donde dibujar: el indicado, o «solo título», o el que tenga menos marcadores."""
    layouts = list(prs.slide_layouts)
    if preferido:
        for l in layouts:
            if l.name == preferido:
                return l
    for clave in ("title only", "solo el título", "solo título", "sólo el título"):
        for l in layouts:
            if l.name.lower() == clave:
                return l
    return min(layouts, key=lambda l: len(l.placeholders))


def pptx_de_guion(meta: dict, slides: list[dict], perfil: dict, carpeta: Path, salida: Path, sufijo: str = "") -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt
    from PIL import Image

    rgb = lambda h: RGBColor.from_string(h.lstrip("#").upper())
    en = meta.get("idioma") == "en"
    c, f = perfil["colores"], perfil["fuentes"]
    fuente_tit, fuente_txt = f.get("pptx_titulo", "Georgia"), f.get("pptx_texto", "Segoe UI")
    plantilla = perfil.get("plantilla_pptx")
    if plantilla:
        prs = Presentation(str(perfil["_carpeta"] / plantilla))
        _borrar_diapositivas(prs)
    else:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)   # 16:9
    W, H = prs.slide_width, prs.slide_height
    px = W / 1920                         # 1 px del escenario HTML en EMU: mismas proporciones que deck.css
    layout = _layout(prs, perfil.get("layout_pptx"))
    M, TOP = int(120 * px), int(90 * px)

    def caja(sl, x, y, w, h, textos, tam, color, fuente, negrita=False, ajustar=True, medio=False):
        from pptx.enum.text import MSO_ANCHOR
        tb = sl.shapes.add_textbox(int(x), int(y), int(w), int(h))
        tf = tb.text_frame
        tf.word_wrap = ajustar
        if medio:
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for k, (t, sangria) in enumerate(textos):
            par = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            par.space_after = Pt(tam * 0.6)
            # **negrita** del guion: misma marca que en el HTML (negrita en color de acento)
            trozos = re.split(r"(\*\*.+?\*\*)", ("•  " if sangria else "") + t)
            for trozo in trozos:
                if not trozo:
                    continue
                marcado = trozo.startswith("**") and trozo.endswith("**")
                r = par.add_run()
                r.text = _plano(trozo)
                r.font.size, r.font.name, r.font.bold = Pt(tam), fuente, negrita or marcado
                r.font.color.rgb = rgb(c["acento"] if marcado and not negrita else color)
        return tb

    def imagen(sl, ruta, x, y, w, h):
        ruta = (carpeta / ruta)
        if ruta.suffix.lower() == ".svg" and ruta.with_suffix(".png").exists():
            ruta = ruta.with_suffix(".png")
        if not ruta.exists() or ruta.suffix.lower() == ".svg":
            caja(sl, x, y, w, h, [(f"[falta la imagen PNG de {ruta.name}]", False)], 16, c["acento_2"], fuente_txt)
            return
        with Image.open(ruta) as im:
            iw, ih = im.size
        esc = min(w / iw, h / ih)
        ancho, alto = int(iw * esc), int(ih * esc)
        sl.shapes.add_picture(str(ruta), int(x + (w - ancho) / 2), int(y + (h - alto) / 2), ancho, alto)

    for d in slides:
        sl = prs.slides.add_slide(layout)
        # Con plantilla, el título va en su marcador (hereda la tipografía institucional); sin plantilla, lo
        # dibujamos con las fuentes del perfil.
        titulo_ph = sl.shapes.title if plantilla and d["tipo"] not in ("portada", "seccion") else None
        if not plantilla:
            sl.background.fill.solid()
            sl.background.fill.fore_color.rgb = rgb(c["fondo"])
        if perfil.get("fondo_imagen"):
            # imagen de fondo a pantalla completa, detrás de todo (justo después de las propiedades del árbol)
            pic = sl.shapes.add_picture(str(perfil["_carpeta"] / perfil["fondo_imagen"]), 0, 0, W, H)
            arbol = sl.shapes._spTree
            arbol.remove(pic._element)
            arbol.insert(2, pic._element)
        for ph in list(sl.placeholders):
            if ph is not titulo_ph:
                ph._element.getparent().remove(ph._element)
        texto = [(p, False) for p in d["parrafos"]] + [(v, True) for v in d["vinetas"]]
        tipo = d["tipo"]
        if tipo == "portada":
            caja(sl, M, H * 0.30, W - 2 * M, 300 * px, [(d["titulo"], False)], 42, c["tinta"], fuente_tit, True)
            if d["subtitulo"]:
                caja(sl, M, H * 0.30 + 330 * px, W - 2 * M, 120 * px, [(d["subtitulo"], False)], 20, c["tinta_suave"], fuente_txt)
            aut = " · ".join(str(x) for x in (meta.get("autor"), meta.get("lugar"), meta.get("fecha")) if x)
            caja(sl, M, H * 0.78, W - 2 * M, 60 * px, [(aut, False)], 14, c["tinta_suave"], fuente_txt)
        elif tipo == "seccion":
            if d["subtitulo"]:
                caja(sl, M, H * 0.36, W - 2 * M, 50 * px, [(d["subtitulo"], False)], 14, c["acento"], fuente_txt)
            caja(sl, M, H * 0.42, W - 2 * M, 200 * px, [(d["titulo"], False)], 38, c["tinta"], fuente_tit, True)
        else:
            alto_tit = 150 * px
            if titulo_ph is not None:
                titulo_ph.text = _plano(d["titulo"])
            else:
                caja(sl, M, TOP, W - 2 * M, alto_tit, [(d["titulo"], False)], 30, c["tinta"], fuente_tit, True)
            y0 = TOP + alto_tit + 30 * px
            y1 = H - (110 if d["fuente"] else 70) * px
            ancho, alto = W - 2 * M, y1 - y0
            figs = list(d["figuras"])
            if d["anim"]:
                auto = f"figuras/estatico_{d['anim']}{sufijo}.png"
                figs = [d["estatico"]] if d["estatico"] else ([auto] if (carpeta / auto).exists() else [])
                if not figs:
                    texto = texto + [("[Animación: ver la versión HTML. Corre «kit.py estaticos» o agrega @estatico.]", False)]
            if tipo == "dato":
                caja(sl, M, y0, 800 * px, alto, [(d["cifra"] or "", False)], 84, c["acento"], fuente_tit, True,
                     ajustar=False, medio=True)
                caja(sl, M + 860 * px, y0, ancho - 860 * px, alto, texto, 19, c["tinta"], fuente_txt, medio=True)
            elif figs and texto:
                lado = 560 * px
                for k, fig in enumerate(figs):
                    w = (ancho - lado - 64 * px) / len(figs)
                    imagen(sl, fig, M + k * w, y0, w - 20 * px, alto)
                caja(sl, W - M - lado, y0, lado, alto, texto, 17, c["tinta"], fuente_txt)
            elif figs:
                w = ancho / len(figs)
                for k, fig in enumerate(figs):
                    imagen(sl, fig, M + k * w, y0, w - (40 * px if len(figs) > 1 else 0), alto)
            else:   # solo texto: más grande y centrado en la altura, igual que en el HTML (deck.css)
                anexo = d["anexo"]
                caja(sl, M, y0, ancho, alto, texto, 14 if anexo else 21, c["tinta"], fuente_txt, medio=not anexo)
        if d["anexo"]:
            caja(sl, M, TOP - 40 * px, 400 * px, 40 * px, [("APPENDIX" if en else "ANEXO", False)], 10, c["tinta_suave"], fuente_txt)
        if d["fuente"]:
            caja(sl, M, H - 95 * px, W - 2 * M - 120 * px, 60 * px, [(("Source: " if en else "Fuente: ") + d["fuente"], False)], 11, c["tinta_suave"], fuente_txt)
        if perfil.get("logo") and not plantilla and tipo != "portada":
            logo = perfil["_carpeta"] / perfil["logo"]
            if logo.exists() and logo.suffix.lower() in (".png", ".jpg", ".jpeg"):
                sl.shapes.add_picture(str(logo), int(W - 56 * px - 200 * px), int(40 * px), height=int(64 * px))
        if d["notas"]:
            sl.notes_slide.notes_text_frame.text = " ".join(_plano(n) for n in d["notas"])
    prs.core_properties.title = str(meta.get("titulo", ""))
    prs.core_properties.author = str(meta.get("autor", ""))
    prs.save(salida)
    print(f"[OK] {salida}  ({len(slides)} diapositivas; diseño de la plantilla: «{layout.name}»)")


# =====================================================================
# Comandos
# =====================================================================

GUION_INICIAL = """---
titulo: {titulo}
autor: Tu nombre
fecha: {fecha}
lugar:
perfil: {perfil}
audiencia: quién la ve, qué sabe y qué no
mensaje: la frase que deben recordar en una semana
---

<!-- Una diapositiva por «## Título». El título es la conclusión, en una oración (principios.md, 4).
     Claves: @tipo portada|seccion|figura|dos|dato|texto  @figura ruta  @fuente texto  @cifra «10 cm/s»
     @anim nombre  @estatico ruta (imagen final de la animación, para PDF y .pptx)  @subtitulo texto
     @pasos (revela párrafos y viñetas uno por uno)  @anexo  @porque (comentario de diseño, tecla E)
     @clase completa (figura a pantalla completa)  @antetitulo texto (línea chica sobre el título)
     Front matter «transicion: fundido» para un fundido corto entre diapositivas.
     EN: the same keys exist in English (@type cover|section|figure|two|number|text, @figure, @source, @number,
     @anim, @static, @subtitle, @steps, @appendix, @why, @class full, @kicker; front matter title, author, date,
     place, profile, transition: fade, motion: respect, lang: en).
     Líneas con «- » son viñetas; con «> » son notas del orador; el resto, párrafos. -->

## {titulo}
@tipo portada
@subtitulo Subtítulo

## Escribe aquí la primera conclusión, como oración
@figura figuras/ejemplo.png
@fuente de dónde sale la evidencia
> Lo que dices mientras se ve esta diapositiva.
"""


def cmd_nueva(args: list[str]) -> None:
    import datetime
    if not args:
        sys.exit("Uso: kit.py nueva <carpeta> [--perfil sobrio] [--titulo \"...\"]")
    carpeta = Path(args[0]).resolve()
    perfil = args[args.index("--perfil") + 1] if "--perfil" in args else "sobrio"
    titulo = args[args.index("--titulo") + 1] if "--titulo" in args else carpeta.name.replace("_", " ")
    origen_perfil = PERFILES / perfil if (PERFILES / perfil).is_dir() else Path(perfil)
    if not (origen_perfil / "perfil.yml").exists():
        sys.exit(f"[ERROR] No encuentro el perfil «{perfil}». Perfiles: {', '.join(p.name for p in PERFILES.iterdir() if p.is_dir())}")
    if (carpeta / "guion.md").exists():
        sys.exit(f"[ERROR] {carpeta} ya tiene un guion.md: no lo piso.")
    carpeta.mkdir(parents=True, exist_ok=True)
    (carpeta / "figuras").mkdir(exist_ok=True)
    (carpeta / "motor").mkdir(exist_ok=True)
    for a in ("deck.css", "deck.js", "anim.js", "figuras.py"):
        shutil.copy2(MOTOR / a, carpeta / "motor" / a)
    shutil.copytree(origen_perfil, carpeta / "perfil", dirs_exist_ok=True)
    (carpeta / "guion.md").write_text(GUION_INICIAL.format(titulo=titulo, perfil=origen_perfil.name,
                                                           fecha=datetime.date.today().isoformat()), encoding="utf-8")
    print(f"[OK] presentación nueva en {carpeta} con el perfil «{origen_perfil.name}».")
    print("     Siguiente: escribe guion.md y corre «kit.py construir» sobre la carpeta.")


def cmd_construir(args: list[str]) -> None:
    if not args:
        sys.exit("Uso: kit.py construir <carpeta> [--html] [--pptx] [--sobrescribir]")
    carpeta = Path(args[0]).resolve()
    hacer_html = "--html" in args or "--pptx" not in args
    hacer_pptx = "--pptx" in args or "--html" not in args
    guion = args[args.index("--guion") + 1] if "--guion" in args else "guion.md"
    suf = _sufijo(guion)
    meta, slides = leer_guion(carpeta / guion)
    perfil = leer_perfil(carpeta / "perfil")
    nombre = carpeta.name
    if hacer_html:
        destino = carpeta / f"deck{suf}.html"
        if destino.exists() and "Generado desde guion" not in destino.read_text(encoding="utf-8")[:300] \
                and "--sobrescribir" not in args:
            destino = carpeta / "borrador.html"
            print("[AVISO] deck.html existe y no lo generó el kit (puede estar hecho a mano): escribo borrador.html.")
        texto = html_de_guion(meta, slides, bool(perfil.get("logo")), carpeta)
        if perfil.get("logo"):
            texto = texto.replace("{logo}", perfil["logo"])
        destino.write_text(texto, encoding="utf-8")
        print(f"[OK] {destino}  ({len(slides)} diapositivas)")
    if hacer_pptx:
        pptx_de_guion(meta, slides, perfil, carpeta, carpeta / f"{nombre}{suf}.pptx", suf)


def _data_uri(ruta: Path) -> str:
    tipo = mimetypes.guess_type(ruta.name)[0] or {".woff2": "font/woff2", ".woff": "font/woff"}.get(ruta.suffix, "application/octet-stream")
    return f"data:{tipo};base64,{base64.b64encode(ruta.read_bytes()).decode()}"


def cmd_empaquetar(args: list[str]) -> None:
    if not args:
        sys.exit("Uso: kit.py empaquetar <deck.html> [salida.html]")
    origen = Path(args[0]).resolve()
    base = origen.parent
    # deck.html -> <carpeta>.html ; deck.en.html -> <carpeta>.en.html
    salida = Path(args[1]).resolve() if len(args) > 1 else base / f"{base.name}{origen.name[4:-5]}.html"
    texto = origen.read_text(encoding="utf-8")
    faltan = []

    def css_inline(m):
        ruta = (base / m.group(1)).resolve()
        if not ruta.exists():
            faltan.append(m.group(1)); return m.group(0)
        css = ruta.read_text(encoding="utf-8")
        def url(mu):
            u = mu.group(1).strip("'\"")
            if u.startswith(("data:", "http:", "https:", "#")):
                return mu.group(0)
            r = (ruta.parent / u).resolve()
            if not r.exists():
                faltan.append(u); return mu.group(0)
            return f"url('{_data_uri(r)}')"
        return "<style>\n" + re.sub(r"url\(([^)]+)\)", url, css) + "\n</style>"

    def js_inline(m):
        ruta = (base / m.group(1)).resolve()
        if not ruta.exists():
            faltan.append(m.group(1)); return m.group(0)
        return "<script>\n" + ruta.read_text(encoding="utf-8").replace("</script", "<\\/script") + "\n</script>"

    def atributo(m):
        attr, u = m.group(1), m.group(2)
        if u.startswith(("data:", "http:", "https:", "#")) or not u:
            return m.group(0)
        r = (base / u).resolve()
        if not r.exists():
            faltan.append(u); return m.group(0)
        return f'{attr}="{_data_uri(r)}"'

    texto = re.sub(r'<link[^>]+rel="stylesheet"[^>]+href="([^"]+)"[^>]*>', css_inline, texto)
    texto = re.sub(r'<script src="([^"]+)"></script>', js_inline, texto)
    texto = re.sub(r'\b(src|data-logo|href)="([^"]+\.(?:png|jpe?g|svg|gif|webp|mp4|webm))"', atributo, texto)
    salida.write_text(texto, encoding="utf-8")
    mb = salida.stat().st_size / 1e6
    print(f"[OK] {salida}  ({mb:.1f} MB, un solo archivo)")
    for f in faltan:
        print(f"[AVISO] no encontré {f}: quedó como referencia externa.")


def _contar_slides(archivo: Path) -> int:
    """Diapositivas del HTML, sin contar los ejemplos que aparecen en comentarios del CSS o JS empaquetados."""
    texto = archivo.read_text(encoding="utf-8")
    texto = re.sub(r"<script.*?</script>|<style.*?</style>|<!--.*?-->", "", texto, flags=re.S)
    return len(re.findall(r'<section class="slide[ "]', texto))


def cmd_pdf(args: list[str]) -> None:
    pos = [a for a in args if not a.startswith("--")]
    if not pos:
        sys.exit("Uso: kit.py pdf <archivo.html> [salida.pdf] [--sobrescribir]")
    archivo = Path(pos[0]).resolve()
    pdf = Path(pos[1]).resolve() if len(pos) > 1 else archivo.with_suffix(".pdf")
    if pdf.exists() and "--sobrescribir" not in args:
        sys.exit(f"[ERROR] Ya existe {pdf.name}. Usa --sobrescribir para reemplazarlo. / Use --overwrite to replace it.")
    edge = navegador()
    n_slides = _contar_slides(archivo)
    perfil_edge = Path(tempfile.gettempdir()) / "kit-presentaciones-edge"
    if pdf.exists():
        pdf.unlink()
    r = subprocess.run([edge, "--headless=new", "--disable-gpu", "--no-pdf-header-footer", "--hide-scrollbars",
                        f"--user-data-dir={perfil_edge}", "--virtual-time-budget=8000", f"--print-to-pdf={pdf}",
                        archivo.as_uri() + "?imprimir"], capture_output=True, text=True, timeout=240)
    if not pdf.exists() or pdf.stat().st_size == 0:
        sys.exit(f"[ERROR] Edge no generó el PDF: {(r.stderr or r.stdout)[-400:]}")
    paginas = len(re.findall(rb"/Type\s*/Page(?!s)", pdf.read_bytes()))
    print(f"[OK] {pdf}  ({paginas} páginas; la presentación tiene {n_slides} diapositivas)")
    if paginas != n_slides:
        print("[AVISO] El número de páginas no coincide con el de diapositivas: abre el PDF y revísalo.")


ROTULOS = {"introducción", "introduccion", "contexto", "metodología", "metodologia", "resultados", "conclusiones",
           "conclusión", "agenda", "índice", "indice", "objetivos", "gracias", "¡gracias!", "preguntas", "resumen",
           "motivación", "motivacion", "antecedentes", "discusión", "discusion", "trabajo futuro", "próximos pasos",
           "introduction", "background", "methods", "methodology", "results", "conclusions", "conclusion", "outline",
           "objectives", "thanks", "thank you", "questions", "summary", "motivation", "discussion", "future work",
           "next steps"}


def cmd_revisar(args: list[str]) -> None:
    if not args:
        sys.exit("Uso: kit.py revisar <deck.html>")
    archivo = Path(args[0]).resolve()
    texto = archivo.read_text(encoding="utf-8")
    secciones = re.findall(r'(<section class="slide[^"]*"[^>]*>)(.*?)</section>', texto, re.S)
    codigo = texto + "".join((archivo.parent / s).read_text(encoding="utf-8")
                             for s in re.findall(r'<script src="([^"]+)"', texto)
                             if not s.startswith("http") and (archivo.parent / s).exists())
    hallazgos = []
    titulos = []
    for i, (abre, cuerpo) in enumerate(secciones, 1):
        clase = re.search(r'class="([^"]+)"', abre).group(1)
        h1 = re.search(r"<h1[^>]*>(.*?)</h1>", cuerpo, re.S)
        t = re.sub(r"<[^>]+>", "", h1.group(1)).strip() if h1 else ""
        titulos.append(t)
        es_especial = any(k in clase for k in ("portada", "seccion", "anexo"))
        if not t:
            hallazgos.append((i, "no tiene título (h1)"))
        elif not es_especial and (t.lower().strip(".:") in ROTULOS or len(t.split()) < 3):
            hallazgos.append((i, f"el título «{t}» es un rótulo, no una conclusión (principio 4)"))
        # En pantalla no cuentan las notas, el «por qué funciona» ni la línea de fuente (letra pequeña aparte).
        sin_notas = re.sub(r'<aside class="(notas|porque)">.*?</aside>|<p class="fuente[^"]*">.*?</p>', "", cuerpo, flags=re.S)
        visible = re.sub(r"<script.*?</script>|<style.*?</style>|<[^>]+>", " ", sin_notas, flags=re.S)
        palabras = len(visible.split())
        if not es_especial and palabras > 75 and "anexo" not in clase:
            hallazgos.append((i, f"tiene {palabras} palabras en pantalla: ¿es una diapositiva o un documento? (principio 5)"))
        if re.search(r"<(img|svg|figure|canvas)\b", cuerpo) and 'class="fuente' not in cuerpo and not es_especial:
            hallazgos.append((i, "tiene figura y no tiene línea de fuente (principio 8)"))
        for src in re.findall(r'<img[^>]+src="([^"]+)"', cuerpo):
            if not src.startswith(("data:", "http")) and not (archivo.parent / src).exists():
                hallazgos.append((i, f"la imagen {src} no existe"))
        anim = re.search(r'data-anim="([^"]+)"', abre)
        if anim and not re.search(rf"\.{re.escape(anim.group(1))}\s*=|\b{re.escape(anim.group(1))}\s*:", codigo):
            hallazgos.append((i, f"pide la animación «{anim.group(1)}» y no encuentro su función"))
        if '<aside class="notas">' not in cuerpo and not es_especial:
            hallazgos.append((i, "no tiene notas del orador (no es error; para ensayar conviene tenerlas)"))
    print(f"Revisión mecánica de {archivo.name}: {len(secciones)} diapositivas\n")
    print("Los títulos, leídos de corrido (¿cuentan la historia solos?):")
    for i, t in enumerate(titulos, 1):
        print(f"  {i:>2}. {t}")
    print()
    if not hallazgos:
        print("[OK] Sin hallazgos mecánicos. Falta la revisión con los ojos: principios.md, «Lista de revisión».")
    for i, h in hallazgos:
        print(f"  diapositiva {i}: {h}")
    print("\nEsto revisa lo que se puede contar. El resto de la lista (principios.md) se revisa mirando.")


def cmd_perfil(args: list[str]) -> None:
    if "--regenerar" in args:
        regenerar_perfil(Path(args[args.index("--regenerar") + 1]).resolve())
        return
    if "--desde" not in args or "--nombre" not in args:
        sys.exit("Uso: kit.py perfil --desde <referencia.pptx> --nombre <nombre> [--destino perfiles]  |  "
                 "kit.py perfil --regenerar <carpeta_perfil>")
    ref = Path(args[args.index("--desde") + 1]).resolve()
    destino = Path(args[args.index("--destino") + 1]).resolve() if "--destino" in args else PERFILES
    if ref.suffix.lower() not in (".pptx", ".potx"):
        sys.exit("[ERROR] Solo leo .pptx o .potx. Para un PDF, una imagen o un HTML: copia perfiles/sobrio, "
                 "edita perfil.yml mirando la referencia y corre «kit.py perfil --regenerar».")
    perfil_desde_pptx(ref, args[args.index("--nombre") + 1], destino)


def cmd_capturas(args: list[str]) -> None:
    """Una imagen por diapositiva del HTML, con Edge sin ventana: para revisar con los ojos (o que un modelo revise)."""
    if len(args) < 2:
        sys.exit("Uso: kit.py capturas <deck.html> <carpeta_png> [--solo 3,5,7]")
    archivo, carpeta = Path(args[0]).resolve(), Path(args[1]).resolve()
    carpeta.mkdir(parents=True, exist_ok=True)
    edge = navegador()
    n = _contar_slides(archivo)
    solo = [int(x) for x in args[args.index("--solo") + 1].split(",")] if "--solo" in args else range(1, n + 1)
    perfil_edge = Path(tempfile.gettempdir()) / "kit-presentaciones-edge"
    for k in solo:
        destino = carpeta / f"diapositiva_{k:02d}.png"
        subprocess.run([edge, "--headless=new", "--disable-gpu", "--hide-scrollbars", f"--user-data-dir={perfil_edge}",
                        "--window-size=1920,1080", "--virtual-time-budget=4000", f"--screenshot={destino}",
                        archivo.as_uri() + f"?imprimir#{k}"], capture_output=True, timeout=120)
        if not destino.exists():
            print(f"[AVISO] no salió la captura de la diapositiva {k}")
    print(f"[OK] {len(list(solo))} capturas en {carpeta} (cada diapositiva en su estado final)")


def cmd_estaticos(args: list[str]) -> None:
    """Fotografía el estado final de cada animación (sin título ni fuente) en figuras/estatico_<nombre>.png.
    El .pptx las usa en lugar de la animación; el PDF no las necesita (imprime el SVG en su estado final)."""
    if not args:
        sys.exit("Uso: kit.py estaticos <carpeta>")
    carpeta = Path(args[0]).resolve()
    suf = _sufijo(args[args.index("--guion") + 1]) if "--guion" in args else ""
    archivo = carpeta / f"deck{suf}.html"
    texto = re.sub(r"<script.*?</script>|<!--.*?-->", "", archivo.read_text(encoding="utf-8"), flags=re.S)
    edge = navegador()
    perfil_edge = Path(tempfile.gettempdir()) / "kit-presentaciones-edge"
    secciones = re.findall(r'<section class="slide[^"]*"([^>]*)>', texto)
    hechas = 0
    for k, attrs in enumerate(secciones, 1):
        m = re.search(r'data-anim="([^"]+)"', attrs)
        if not m:
            continue
        destino = carpeta / "figuras" / f"estatico_{m.group(1)}{suf}.png"
        destino.parent.mkdir(exist_ok=True)
        subprocess.run([edge, "--headless=new", "--disable-gpu", "--hide-scrollbars", f"--user-data-dir={perfil_edge}",
                        "--window-size=1920,1080", "--virtual-time-budget=5000", f"--screenshot={destino}",
                        archivo.as_uri() + f"?imprimir&solo-lienzo#{k}"], capture_output=True, timeout=120)
        hechas += destino.exists()
        print(f"  diapositiva {k}: {destino.name}" if destino.exists() else f"[AVISO] falló la diapositiva {k}")
    print(f"[OK] {hechas} estáticos en {carpeta / 'figuras'}. Vuelve a correr «construir --pptx» para usarlos.")


def cmd_renderizar(args: list[str]) -> None:
    if len(args) < 2:
        sys.exit("Uso: kit.py renderizar <archivo.pptx> <carpeta_png>")
    archivo, carpeta = Path(args[0]).resolve(), Path(args[1]).resolve()
    carpeta.mkdir(parents=True, exist_ok=True)
    ps = f"""
$pp = New-Object -ComObject PowerPoint.Application
$pres = $pp.Presentations.Open('{archivo}', $true, $false, $false)
$i = 1
foreach ($s in $pres.Slides) {{ $s.Export((Join-Path '{carpeta}' ('diapositiva_{{0:D2}}.png' -f $i)), 'PNG', 1600, 900); $i++ }}
$pres.Close(); $pp.Quit()
"""
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)
    print(f"[OK] imágenes en {carpeta}")


COMANDOS = {"nueva": cmd_nueva, "construir": cmd_construir, "empaquetar": cmd_empaquetar, "pdf": cmd_pdf,
            "revisar": cmd_revisar, "perfil": cmd_perfil, "renderizar": cmd_renderizar, "capturas": cmd_capturas, "estaticos": cmd_estaticos}

ALIAS = {"new": "nueva", "build": "construir", "package": "empaquetar", "review": "revisar", "profile": "perfil",
         "render": "renderizar", "shots": "capturas", "statics": "estaticos"}
OPCIONES_EN = {"--profile": "--perfil", "--title": "--titulo", "--overwrite": "--sobrescribir", "--only": "--solo",
               "--from": "--desde", "--name": "--nombre", "--dest": "--destino", "--regenerate": "--regenerar",
               "--script": "--guion"}

if __name__ == "__main__":
    comando = ALIAS.get(sys.argv[1], sys.argv[1]) if len(sys.argv) > 1 else ""
    if comando not in COMANDOS:
        print(__doc__)
        sys.exit(2)
    COMANDOS[comando]([OPCIONES_EN.get(a, a) for a in sys.argv[2:]])
