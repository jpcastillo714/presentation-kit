"""presentation-kit: one script for all the mechanics.

Design does not live here: it lives in the profile (colors, fonts, logo, template) and in the script (what each slide
says). This tool turns the script into HTML and .pptx, packages, exports and reviews.

Commands (Spanish aliases in parentheses):

  python engine/kit.py new (nueva) <folder> [--profile sober] [--title "..."] [--lang es]
      Scaffolds a deck: copies the engine and the profile into <folder> and writes a starter script.md.
  python engine/kit.py build (construir) <folder> [--html] [--pptx] [--script script.es.md]
      Script -> deck.html and <folder>.pptx (script.es.md -> deck.es.html and <folder>.es.pptx). Both by default.
  python engine/kit.py package (empaquetar) <deck.html> [out.html]
      Single self-contained HTML file (CSS, JS, images, fonts inside): works offline and can be sent around.
  python engine/kit.py pdf <file.html> [out.pdf] [--overwrite]
      PDF via Chrome or Edge, one slide per page, animations in their final state.
  python engine/kit.py profile (perfil) --from <ref.pptx> --name <name> [--dest profiles]
      Profile from a reference deck: theme and actually-used colors, fonts, logos, background, template.
  python engine/kit.py profile --regenerate <profile_dir>
      Rewrites profile.css and figures.mplstyle from profile.yml (after editing it by hand).
  python engine/kit.py review (revisar) <deck.html>
      Mechanical review against PRINCIPLES.md: label-titles, figures without source, too much text, missing
      images or animations.
  python engine/kit.py shots (capturas) <deck.html> <png_dir> [--only 3,5]
      One PNG per slide (final state), to review by eye.
  python engine/kit.py statics (estaticos) <folder> [--script script.es.md]
      Final frame of each animation (figures/static_<name>.png), used by the .pptx.
  python engine/kit.py demo <demo.yml>
      Records a demo video (MP4 + poster) of a web interface from the REAL running app: zoom where the action is,
      visible cursor, captions. Use it in a slide with «@video figures/<name>.mp4».
  python engine/kit.py guide (guia) new <folder> [--title "..."] [--lang es]   |   kit.py guide <folder>
      Interactive study guide: pages/*.html → one self-contained HTML per page plus a map (index.html), with
      simulators, «predict, then look» questions and an exercise mode. See examples/study-guide.
  python engine/kit.py render (renderizar) <file.pptx> <png_dir>
      .pptx -> PNG with PowerPoint (Windows only), to review it by eye.

Spanish scripts work too: script.es.md (or lang: es) produces a Spanish deck, and the Spanish keys of the script
(@figura, @fuente, @cifra, ...) and of the front matter (titulo, perfil, ...) are accepted as aliases.
"""
from __future__ import annotations

import base64
import html
import mimetypes
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import yaml

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

KIT = Path(__file__).resolve().parent.parent          # kit root
ENGINE = KIT / "engine"
PROFILES = KIT / "profiles"
BROWSERS = [r"C:\Program Files\Google\Chrome\Application\chrome.exe",
            r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
            r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
            r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
            "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
            "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
            "/Applications/Chromium.app/Contents/MacOS/Chromium"]


def browser() -> str:
    """Chrome, Edge or Chromium in headless mode (PDF, screenshots). Set KIT_BROWSER to force a specific binary."""
    if os.environ.get("KIT_BROWSER"):
        return os.environ["KIT_BROWSER"]
    for path in BROWSERS:
        if Path(path).exists():
            return path
    for name in ("google-chrome", "google-chrome-stable", "chromium", "chromium-browser", "microsoft-edge",
                 "msedge", "chrome"):
        if shutil.which(name):
            return shutil.which(name)
    sys.exit("[ERROR] No Chrome, Edge or Chromium found. Install one or set KIT_BROWSER.")


def script_suffix(script_name: str) -> str:
    """script.md -> '' ; script.es.md -> '.es' (guion*.md is accepted too)."""
    m = re.match(r"(?:script|guion)(.*)\.md$", script_name)
    return m.group(1) if m else ""


# =====================================================================
# Profiles
# =====================================================================

def read_profile(folder: Path) -> dict:
    data = yaml.safe_load((folder / "profile.yml").read_text(encoding="utf-8"))
    data["_dir"] = folder
    return data


def profile_css(p: dict) -> str:
    c, f, s = p["colors"], p["fonts"], p.get("sizes", {})
    faces = []
    fonts_dir = p["_dir"] / "fonts"
    if fonts_dir.is_dir():
        for file in sorted(fonts_dir.iterdir()):
            if file.suffix.lower() not in (".woff2", ".woff", ".ttf", ".otf"):
                continue
            # Naming convention: Family-Weight[-Italic].ext, e.g. SourceSerif4-600.woff2
            m = re.match(r"(.+?)-(\d{3})(-italic)?$", file.stem, re.I)
            family, weight, it = (m.group(1), m.group(2), m.group(3)) if m else (file.stem, "400", None)
            faces.append(f"@font-face {{ font-family: '{family}'; src: url('fonts/{file.name}'); "
                         f"font-weight: {weight}; font-style: {'italic' if it else 'normal'}; }}")
    v = {
        "--background": c["background"], "--ink": c["ink"], "--ink-soft": c["ink_soft"], "--accent": c["accent"],
        "--accent-2": c["accent_2"], "--line": c["line"], "--frame": c.get("frame", "#111111"),
        "--font-title": f["title"], "--font-text": f["text"], "--font-mono": f.get("mono", "Consolas, monospace"),
        "--title-size": f"{s.get('title', 60)}px", "--text-size": f"{s.get('text', 34)}px",
        "--title-weight": str(s.get("title_weight", 600)),
    }
    for i, col in enumerate(c.get("series", []), 1):
        v[f"--series-{i}"] = col
    body = "\n".join(f"  {k}: {val};" for k, val in v.items())
    extra = (p.get("css_extra") or "").strip()
    if p.get("background_image"):
        extra = f".slide {{ background-image: url('{p['background_image']}'); background-size: cover; }}\n" + extra
    return (f"/* Profile «{p['name']}», generated from profile.yml by kit.py. Edit the .yml, not this file. */\n"
            + "\n".join(faces) + ("\n" if faces else "") + f":root {{\n{body}\n}}\n" + (extra + "\n" if extra else ""))


def profile_mplstyle(p: dict) -> str:
    c, f = p["colors"], p["fonts"]
    hexes = ", ".join(f"'{s.lstrip('#')}'" for s in c.get("series", [c["accent"], c["accent_2"]]))
    families = ", ".join(f.get("figures", ["DejaVu Sans"]))
    q = lambda s: s.lstrip("#")
    # Sizes meant for a figure that fills the body of a 1920×1080 slide (principle 14).
    return f"""# Figure style of profile «{p['name']}», generated from profile.yml by kit.py.
figure.figsize : 12, 5.4
figure.dpi : 100
savefig.dpi : 160
savefig.bbox : tight
savefig.facecolor : {q(c['background'])}
figure.facecolor : {q(c['background'])}
axes.facecolor : {q(c['background'])}
axes.edgecolor : {q(c['ink_soft'])}
axes.labelcolor : {q(c['ink'])}
axes.prop_cycle : cycler('color', [{hexes}])
axes.spines.top : False
axes.spines.right : False
axes.grid : False
axes.titlesize : 18
axes.labelsize : 17
axes.titlelocation : left
xtick.color : {q(c['ink_soft'])}
ytick.color : {q(c['ink_soft'])}
xtick.labelsize : 15
ytick.labelsize : 15
text.color : {q(c['ink'])}
font.family : sans-serif
font.sans-serif : {families}, DejaVu Sans
font.size : 16
legend.frameon : False
legend.fontsize : 15
lines.linewidth : 2.4
lines.markersize : 7
image.cmap : cividis
svg.fonttype : none
"""


def regenerate_profile(folder: Path) -> None:
    p = read_profile(folder)
    (folder / "profile.css").write_text(profile_css(p), encoding="utf-8")
    (folder / "figures.mplstyle").write_text(profile_mplstyle(p), encoding="utf-8")
    print(f"[OK] profile «{p['name']}»: profile.css and figures.mplstyle regenerated in {folder}")


NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "rel": "http://schemas.openxmlformats.org/package/2006/relationships"}


def _theme_color(el) -> str | None:
    if el is None:
        return None
    s = el.find("a:srgbClr", NS)
    if s is not None:
        return "#" + s.get("val").lower()
    s = el.find("a:sysClr", NS)
    if s is not None and s.get("lastClr"):
        return "#" + s.get("lastClr").lower()
    return None


def _luminance(hexa: str) -> float:
    r, g, b = (int(hexa[i:i + 2], 16) / 255 for i in (1, 3, 5))
    lin = lambda x: x / 12.92 if x <= 0.03928 else ((x + 0.055) / 1.055) ** 2.4
    return 0.2126 * lin(r) + 0.7152 * lin(g) + 0.0722 * lin(b)


def contrast(a: str, b: str) -> float:
    la, lb = sorted((_luminance(a), _luminance(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


def _blend(a: str, b: str, t: float) -> str:
    ca = [int(a[i:i + 2], 16) for i in (1, 3, 5)]
    cb = [int(b[i:i + 2], 16) for i in (1, 3, 5)]
    return "#" + "".join(f"{round(x + (y - x) * t):02x}" for x, y in zip(ca, cb))


def profile_from_pptx(ref: Path, name: str, dest: Path) -> Path:
    """Reads the theme of a reference deck and builds a profile. Whatever cannot be read is reported."""
    folder = dest / name
    if folder.exists():
        sys.exit(f"[ERROR] Profile {folder} already exists. Pick another name or remove it yourself.")
    warnings = []
    with zipfile.ZipFile(ref) as z:
        names = z.namelist()
        themes = sorted(n for n in names if re.match(r"ppt/theme/theme\d+\.xml$", n))
        if not themes:
            sys.exit("[ERROR] The reference has no theme: is it a .pptx?")
        theme = ET.fromstring(z.read(themes[0]))
        scheme = theme.find(".//a:clrScheme", NS)
        col = {child.tag.split("}")[1]: _theme_color(child) for child in scheme} if scheme is not None else {}
        major = theme.find(".//a:majorFont/a:latin", NS)
        minor = theme.find(".//a:minorFont/a:latin", NS)
        f_title = major.get("typeface") if major is not None else "Georgia"
        f_text = minor.get("typeface") if minor is not None else "Segoe UI"
        # Images repeated on many slides or placed in the master: probably logos (or backgrounds).
        use: dict[str, int] = {}
        slide_rels = [n for n in names if re.match(r"ppt/slides/_rels/slide\d+\.xml\.rels$", n)]
        master_rels = [n for n in names if re.match(r"ppt/(slideMasters|slideLayouts)/_rels/.+\.rels$", n)]
        for n in slide_rels + master_rels:
            for r in ET.fromstring(z.read(n)).findall("rel:Relationship", NS):
                t = r.get("Target", "")
                if "/media/" in t or t.startswith("../media/"):
                    key = "ppt/media/" + t.split("media/")[1]
                    use[key] = use.get(key, 0) + (100 if n in master_rels else 1)
        n_slides = max(1, len(slide_rels))
        # Colors the slides actually use (fills and text): the theme may be Office's default even when the deck uses
        # another palette (common in decks generated by code).
        used: dict[str, int] = {}
        for n in names:
            if re.match(r"ppt/(slides/slide|slideMasters/slideMaster|slideLayouts/slideLayout)\d+\.xml$", n):
                for val in re.findall(rb'<a:srgbClr val="([0-9A-Fa-f]{6})"', z.read(n)):
                    k = "#" + val.decode().lower()
                    used[k] = used.get(k, 0) + 1
        used_colors = [k for k, _ in sorted(used.items(), key=lambda x: -x[1])[:10]]
        candidates = [m for m, k in sorted(use.items(), key=lambda x: -x[1]) if k >= 100 or k >= 0.5 * n_slides]
        folder.mkdir(parents=True)
        logos, backgrounds = [], []
        for m in candidates[:6]:
            if m in names and Path(m).suffix.lower() in (".png", ".jpg", ".jpeg", ".svg", ".gif", ".emf", ".wmf"):
                data = z.read(m)
                is_background = False
                try:   # a large image shaped like the slide is a background, not a logo
                    from io import BytesIO
                    from PIL import Image
                    with Image.open(BytesIO(data)) as im:
                        w_img, h_img = im.size
                    is_background = w_img >= 900 and 1.2 <= w_img / h_img <= 1.9
                except Exception:
                    pass
                sub = "backgrounds" if is_background else "logos"
                (folder / sub).mkdir(exist_ok=True)
                (folder / sub / Path(m).name).write_bytes(data)
                (backgrounds if is_background else logos).append(Path(m).name)
    background = col.get("lt1") or "#ffffff"
    ink = col.get("dk1") or "#1a1a1a"
    background_image = None
    if backgrounds:
        # The background is an image placed on the slides (not in the master): it becomes the profile background, and
        # the background color and ink follow its tone so text keeps its contrast.
        from PIL import Image
        with Image.open(folder / "backgrounds" / backgrounds[0]) as im:
            r_, g_, b_ = im.convert("RGB").resize((1, 1), Image.Resampling.BOX).getpixel((0, 0))
        background = f"#{r_:02x}{g_:02x}{b_:02x}"
        background_image = f"backgrounds/{backgrounds[0]}"
        if _luminance(background) < 0.2:
            lights = [c_ for c_ in used_colors if _luminance(c_) > 0.7]
            ink = lights[0] if lights else "#f2f2f2"
            warnings.append(f"The background is a dark image ({background}): ink switched to {ink}. matplotlib "
                            "figures made with this profile will have a dark background.")
        col["lt2"] = None
    ink_soft = col.get("dk2") or _blend(ink, background, 0.35)
    if contrast(ink_soft, background) < 4.5:
        ink_soft = _blend(ink, background, 0.35)
        warnings.append("The theme's dk2 does not contrast enough with the background; ink_soft was blended.")
    accents = [col.get(f"accent{i}") for i in range(1, 7) if col.get(f"accent{i}")]
    if not accents:
        accents = ["#1f5f8b", "#b4532a"]
        warnings.append("The theme has no accents; the sober profile's were used.")
    for a in accents[:2]:
        if contrast(a, background) < 3:
            warnings.append(f"Accent {a} has low contrast with background {background} (< 3:1): use it only in large areas.")
    line = col.get("lt2") or _blend(background, ink, 0.15)
    if col.get("accent1") == "#4472c4" and col.get("accent2") == "#ed7d31":
        warnings.append("The theme is Office's default: the real palette is most likely in used_colors. Assign "
                        "background, ink and accents from there and run «kit.py profile --regenerate».")
    shutil.copy2(ref, folder / "template.pptx")
    data = {
        "name": name,
        "description": f"Extracted from {ref.name}. Check colors and logo before using it.",
        "origin": ref.name,
        "colors": {"background": background, "ink": ink, "ink_soft": ink_soft, "accent": accents[0],
                   "accent_2": accents[1] if len(accents) > 1 else _blend(accents[0], ink, 0.4),
                   "line": line, "frame": "#111111", "series": accents[:5]},
        "fonts": {"title": f"'{f_title}', Georgia, serif", "text": f"'{f_text}', 'Segoe UI', system-ui, sans-serif",
                  "mono": "Consolas, monospace", "pptx_title": f_title, "pptx_text": f_text,
                  "figures": [f_text, "Segoe UI"]},
        "sizes": {"title": 60, "text": 34, "title_weight": 600},
        "logo": f"logos/{logos[0]}" if logos else None,
        "background_image": background_image,
        "pptx_template": "template.pptx",
        "pptx_layout": None,
        "used_colors": used_colors,
        "warnings": warnings,
    }
    (folder / "profile.yml").write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False), encoding="utf-8")
    regenerate_profile(folder)
    print(f"[OK] profile «{name}» in {folder}")
    print(f"     theme colors: {', '.join(f'{k}={v}' for k, v in col.items() if v)}")
    print(f"     theme fonts: titles «{f_title}», text «{f_text}»")
    print(f"     logo candidates: {', '.join(logos) or 'none'}")
    if backgrounds:
        print(f"     slide background image (background_image): {', '.join(backgrounds)}")
    print(f"     most used colors on the slides: {', '.join(used_colors) or 'none explicit'}")
    for w in warnings:
        print(f"[WARNING] {w}")
    print("     Look at the profile before using it: the theme says which colors it offers, not which ones the deck uses.")
    return folder


# =====================================================================
# Script
# =====================================================================

KEYS_ES = {"tipo": "type", "figura": "figure", "fuente": "source", "cifra": "number", "estatico": "static",
           "subtitulo": "subtitle", "pasos": "steps", "anexo": "appendix", "porque": "why", "clase": "class",
           "antetitulo": "kicker"}
TYPES_ES = {"portada": "cover", "seccion": "section", "figura": "figure", "dos": "two", "dato": "number",
            "texto": "text"}
META_ES = {"titulo": "title", "autor": "author", "fecha": "date", "lugar": "place", "perfil": "profile",
           "transicion": "transition", "movimiento": "motion", "idioma": "lang", "language": "lang",
           "audiencia": "audience", "mensaje": "message"}
VALUES_ES = {"fundido": "fade", "respetar": "respect", "completa": "full"}


def read_script(path: Path) -> tuple[dict, list[dict]]:
    """Script: YAML front matter + one slide per «## Title». See the starter script that «new» writes."""
    text = path.read_text(encoding="utf-8")
    meta = {}
    m = re.match(r"^---\n(.*?)\n---\n", text, re.S)
    if m:
        meta = yaml.safe_load(m.group(1)) or {}
        meta = {META_ES.get(k, k): (VALUES_ES.get(v, v) if isinstance(v, str) else v) for k, v in meta.items()}
        text = text[m.end():]
    meta.setdefault("lang", script_suffix(path.name).lstrip(".") or "en")
    slides, sl = [], None
    for line in text.splitlines():
        if line.startswith("## "):
            sl = {"title": line[3:].strip(), "type": None, "figures": [], "bullets": [], "paragraphs": [],
                  "notes": [], "source": None, "anim": None, "static": None, "number": None, "appendix": False,
                  "steps": False, "subtitle": None, "why": None, "class": None, "kicker": None, "video": None}
            slides.append(sl)
            continue
        if sl is None or not line.strip() or line.startswith("<!--"):
            continue
        s = line.strip()
        if s.startswith("@"):
            key, _, value = s[1:].partition(" ")
            value = value.strip()
            key = KEYS_ES.get(key, key)
            if key == "type":
                value = TYPES_ES.get(value, value)
            if key == "class":
                value = " ".join(VALUES_ES.get(v, v) for v in value.split())
            if key == "figure":
                sl["figures"].append(value)
            elif key in ("appendix", "steps"):
                sl[key] = int(value) if value.isdigit() else True   # «@steps 3»: levels of an animation
            elif key in sl:
                sl[key] = value
            else:
                print(f"[WARNING] unknown key @{key} in «{sl['title']}»")
        elif s.startswith(">"):
            sl["notes"].append(s.lstrip("> ").strip())
        elif s.startswith("- "):
            sl["bullets"].append(s[2:].strip())
        else:
            sl["paragraphs"].append(s)
    for sl in slides:
        if not sl["type"]:
            sl["type"] = ("number" if sl["number"] else "two" if len(sl["figures"]) > 1
                          else "figure" if sl["figures"] or sl["anim"] or sl["video"] else "text")
    return meta, slides


def _md(s: str) -> str:
    """Minimal inline markdown: **bold**, *emphasis*, `code`, ^superscript^."""
    s = html.escape(s, quote=False)
    s = re.sub(r"\^(.+?)\^", r"<sup>\1</sup>", s)
    s = re.sub(r"\*\*(.+?)\*\*", r'<b class="accent">\1</b>', s)
    s = re.sub(r"\*(.+?)\*", r"<em>\1</em>", s)
    return re.sub(r"`(.+?)`", r"<code>\1</code>", s)


_SUP = str.maketrans("0123456789−-+", "⁰¹²³⁴⁵⁶⁷⁸⁹⁻⁻⁺")


def _plain(s: str) -> str:
    """Text for the .pptx: no markdown marks; ^−10^ becomes superscript characters."""
    s = re.sub(r"\^(.+?)\^", lambda m: m.group(1).translate(_SUP), s)
    return re.sub(r"\*\*|\*|`", "", s)


def html_from_script(meta: dict, slides: list[dict], logo: str | None, folder: Path | None = None) -> str:
    """deck.html is regenerated from the script. What is specific to a deck lives apart and is linked if present:
    styles.css (design tweaks) and animations.js (window.animations.<name>)."""
    es = meta.get("lang") == "es"
    parts = []
    for i, sl in enumerate(slides, 1):
        t = sl["type"]
        classes = ["slide", t] + (sl["class"].split() if sl["class"] else [])
        if sl["appendix"]:
            classes.append("appendix")
        attrs = f' data-anim="{sl["anim"]}"' if sl["anim"] else ""
        if type(sl["steps"]) is int:      # bool does not count: True and False are ints in Python too
            attrs += f' data-steps="{sl["steps"]}"'
        step = (lambda k: f' data-step="{k}"') if sl["steps"] is True else (lambda k: "")
        h = [f'<section class="{" ".join(classes)}"{attrs}>']
        if sl["appendix"]:
            h.append(f'  <div class="appendix-label">{"Anexo" if es else "Appendix"}</div>')
        if t == "cover":
            if sl["anim"]:
                h.append(f'  <div class="content"><figure><div class="scene" id="scene-{i}"></div></figure></div>')
            h.append(f"  <h1>{_md(sl['title'])}</h1>")
            if sl["subtitle"]:
                h.append(f'  <p class="subtitle">{_md(sl["subtitle"])}</p>')
            byline = " · ".join(str(x) for x in (meta.get("author"), meta.get("place"), meta.get("date")) if x)
            h.append(f'  <p class="byline">{html.escape(byline)}</p>')
        elif t == "section":
            if sl["subtitle"]:
                h.append(f'  <div class="section-number">{_md(sl["subtitle"])}</div>')
            h.append(f"  <h1>{_md(sl['title'])}</h1>")
        else:
            if sl["kicker"]:
                h.append(f'  <p class="kicker">{_md(sl["kicker"])}</p>')
            h.append(f"  <h1>{_md(sl['title'])}</h1>")
            text = [f"<p{step(k)}>{_md(p)}</p>" for k, p in enumerate(sl["paragraphs"], 1)]
            if sl["bullets"]:
                k0 = len(text)
                text.append("<ul>" + "".join(f"<li{step(k0 + k)}>{_md(v)}</li>" for k, v in enumerate(sl["bullets"], 1)) + "</ul>")
            figs = [f'<figure><img src="{html.escape(f)}" alt=""></figure>' for f in sl["figures"]]
            if sl["anim"]:
                # On screen, the animated scene; in the PDF, the static image (@static) if there is one.
                static = (f'<img class="print-only" src="{html.escape(sl["static"])}" alt="">' if sl["static"] else "")
                figs = [f'<figure><div class="scene{" screen-only" if static else ""}" id="scene-{i}"><!-- animation '
                        f'«{sl["anim"]}»: draw it in animations.js --></div>{static}</figure>']
            if sl["video"]:
                # Demo video (kit.py demo): plays when the slide opens; the PDF shows its poster image.
                v = html.escape(sl["video"])
                poster = html.escape(Path(sl["video"]).with_suffix(".png").as_posix())
                figs = [f'<figure><video class="demo-video screen-only" src="{v}" poster="{poster}" muted playsinline '
                        f'preload="auto"></video><img class="print-only" src="{poster}" alt=""></figure>']
            if t == "number":
                block = [f'<div class="side"><div class="big-number">{_md(sl["number"])}</div></div>',
                         f'<div>{"".join(text)}</div>']
                h.append('  <div class="content">' + "".join(block) + "</div>")
            elif t in ("figure", "two") and text and "full" not in classes:
                h.append('  <div class="content">' + "".join(figs) + f'<div class="side">{"".join(text)}</div></div>')
            elif t in ("figure", "two"):
                h.append(f'  <div class="content{" two" if t == "two" else ""}">' + "".join(figs) + "</div>")
            else:
                h.append('  <div class="content column top">' + "".join(text) + "</div>")
        if sl["source"]:
            h.append(f'  <p class="source">{_md(sl["source"])}</p>')
        if sl["notes"]:
            h.append('  <aside class="notes">' + " ".join(_md(n) for n in sl["notes"]) + "</aside>")
        if sl["why"]:
            h.append(f'  <aside class="why">{_md(sl["why"])}</aside>')
        h.append("</section>")
        parts.append("\n".join(h))
    body = f' data-logo="profile/{logo}"' if logo else ""
    if meta.get("transition") == "fade":
        body += ' class="fade"'
    # Animations are seen by the room, not by the presenter's computer: decks animate even if the OS asks to reduce
    # motion. «motion: respect» in the script hands that decision back to the OS.
    body += f' data-motion="{"respect" if meta.get("motion") == "respect" else "always"}"'
    styles = '\n<link rel="stylesheet" href="styles.css">' if folder and (folder / "styles.css").exists() else ""
    anim = ('\n<script src="engine/anim.js"></script>\n<script src="animations.js"></script>'
            if folder and (folder / "animations.js").exists() else "")
    return f"""<!doctype html>
<!-- Generated from the script by engine/kit.py. Do not edit by hand: edit the script and rebuild.
     Design tweaks go in styles.css; animations in animations.js (linked automatically if present). -->
<html lang="{meta.get('lang', 'en')}">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html.escape(str(meta.get('title', 'Presentation')))}</title>
<link rel="stylesheet" href="engine/deck.css">
<link rel="stylesheet" href="profile/profile.css">{styles}
</head>
<body{body}>
<div class="stage">

{chr(10).join(parts)}

</div>
<script>window.animations = window.animations || {{}};</script>{anim}
<script src="engine/deck.js"></script>
</body>
</html>
"""


# =====================================================================
# .pptx
# =====================================================================

def _delete_slides(prs) -> None:
    lst = prs.slides._sldIdLst
    for sid in list(lst):
        prs.part.drop_rel(sid.rId)
        lst.remove(sid)


def _layout(prs, preferred: str | None):
    """The template layout to draw on: the one named, or «title only», or the one with the fewest placeholders."""
    layouts = list(prs.slide_layouts)
    if preferred:
        for l in layouts:
            if l.name == preferred:
                return l
    for key in ("title only", "solo el título", "solo título", "sólo el título"):
        for l in layouts:
            if l.name.lower() == key:
                return l
    return min(layouts, key=lambda l: len(l.placeholders))


def pptx_from_script(meta: dict, slides: list[dict], profile: dict, folder: Path, out: Path, suffix: str = "") -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt
    from PIL import Image

    rgb = lambda h: RGBColor.from_string(h.lstrip("#").upper())
    es = meta.get("lang") == "es"
    c, f = profile["colors"], profile["fonts"]
    font_title, font_text = f.get("pptx_title", "Georgia"), f.get("pptx_text", "Segoe UI")
    template = profile.get("pptx_template")
    if template:
        prs = Presentation(str(profile["_dir"] / template))
        _delete_slides(prs)
    else:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)   # 16:9
    W, H = prs.slide_width, prs.slide_height
    px = W / 1920                         # 1 px of the HTML stage in EMU: same proportions as deck.css
    layout = _layout(prs, profile.get("pptx_layout"))
    M, TOP = int(120 * px), int(90 * px)

    def box(sl, x, y, w, h, texts, size, color, font, bold=False, wrap=True, middle=False):
        from pptx.enum.text import MSO_ANCHOR
        tb = sl.shapes.add_textbox(int(x), int(y), int(w), int(h))
        tf = tb.text_frame
        tf.word_wrap = wrap
        if middle:
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for k, (t, bullet) in enumerate(texts):
            par = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            par.space_after = Pt(size * 0.6)
            # **bold** in the script: same mark as in the HTML (bold in the accent color)
            for chunk in re.split(r"(\*\*.+?\*\*)", ("•  " if bullet else "") + t):
                if not chunk:
                    continue
                marked = chunk.startswith("**") and chunk.endswith("**")
                r = par.add_run()
                r.text = _plain(chunk)
                r.font.size, r.font.name, r.font.bold = Pt(size), font, bold or marked
                r.font.color.rgb = rgb(c["accent"] if marked and not bold else color)
        return tb

    def image(sl, path, x, y, w, h):
        path = folder / path
        if path.suffix.lower() == ".svg" and path.with_suffix(".png").exists():
            path = path.with_suffix(".png")
        if not path.exists() or path.suffix.lower() == ".svg":
            box(sl, x, y, w, h, [(f"[missing PNG for {path.name}]", False)], 16, c["accent_2"], font_text)
            return
        with Image.open(path) as im:
            iw, ih = im.size
        k = min(w / iw, h / ih)
        width, height = int(iw * k), int(ih * k)
        sl.shapes.add_picture(str(path), int(x + (w - width) / 2), int(y + (h - height) / 2), width, height)

    for d in slides:
        sl = prs.slides.add_slide(layout)
        # With a template, the title goes in its placeholder (inherits the institutional typography); without one, it
        # is drawn with the profile's fonts.
        title_ph = sl.shapes.title if template and d["type"] not in ("cover", "section") else None
        if not template:
            sl.background.fill.solid()
            sl.background.fill.fore_color.rgb = rgb(c["background"])
        if profile.get("background_image"):
            # full-slide background image, behind everything (right after the tree properties)
            pic = sl.shapes.add_picture(str(profile["_dir"] / profile["background_image"]), 0, 0, W, H)
            tree = sl.shapes._spTree
            tree.remove(pic._element)
            tree.insert(2, pic._element)
        # A full-bleed animated slide («@class full»): its final frame (statics, without title or source) becomes the
        # slide background and the text goes on top as editable text, like in the HTML. Placing that frame inside the
        # body, under the title, looked like a photo pasted on the slide.
        full_static = None
        if d["anim"] and "full" in (d["class"] or "").split():
            cand = folder / (d["static"] or f"figures/static_{d['anim']}{suffix}.png")
            if cand.exists():
                full_static = cand
                pic = sl.shapes.add_picture(str(cand), 0, 0, W, H)
                tree = sl.shapes._spTree
                tree.remove(pic._element)
                tree.insert(2 + (1 if profile.get("background_image") else 0), pic._element)
        for ph in list(sl.placeholders):
            if ph is not title_ph:
                ph._element.getparent().remove(ph._element)
        text = [(p, False) for p in d["paragraphs"]] + [(v, True) for v in d["bullets"]]
        kind = d["type"]
        if kind == "cover":
            # With a full-bleed animated cover the text sits in the upper half, clear of the figure (as in the HTML).
            if full_static:
                lines = max(1, -(-len(d["title"]) // 36))  # about 36 characters per line at 42 pt
                y_t = H * 0.12
                y_s = y_t + (90 + 100 * lines) * px
                y_b = y_s + 120 * px
            else:
                y_t, y_s, y_b = H * 0.30, H * 0.30 + 330 * px, H * 0.78
            box(sl, M, y_t, W - 2 * M, 300 * px, [(d["title"], False)], 42, c["ink"], font_title, True)
            if d["subtitle"]:
                box(sl, M, y_s, W - 2 * M, 120 * px, [(d["subtitle"], False)], 20,
                    c["accent"] if full_static else c["ink_soft"], font_text)
            byline = " · ".join(str(x) for x in (meta.get("author"), meta.get("place"), meta.get("date")) if x)
            box(sl, M, y_b, W - 2 * M, 60 * px, [(byline, False)], 14, c["ink_soft"], font_text)
        elif kind == "section":
            if d["subtitle"]:
                box(sl, M, H * 0.36, W - 2 * M, 50 * px, [(d["subtitle"], False)], 14, c["accent"], font_text)
            box(sl, M, H * 0.42, W - 2 * M, 200 * px, [(d["title"], False)], 38, c["ink"], font_title, True)
        else:
            title_h = 150 * px
            top = TOP
            if d["kicker"]:                     # small line above the title, as in the HTML
                box(sl, M, TOP - 20 * px, W - 2 * M, 40 * px, [(_plain(d["kicker"]).upper(), False)], 11, c["accent"], font_text)
                top = TOP + 30 * px
            if title_ph is not None:
                title_ph.text = _plain(d["title"])
            else:
                box(sl, M, top, W - 2 * M, title_h, [(d["title"], False)], 30, c["ink"], font_title, True)
            y0 = TOP + title_h + 30 * px
            y1 = H - (110 if d["source"] else 70) * px
            width, height = W - 2 * M, y1 - y0
            figs = list(d["figures"])
            if d["video"]:
                movie = folder / d["video"]
                poster = movie.with_suffix(".png")
                if movie.exists():
                    k = min(width / 1920, height / 1080)
                    vw_, vh_ = int(1920 * k), int(1080 * k)
                    sl.shapes.add_movie(str(movie), int(M + (width - vw_) / 2), int(y0 + (height - vh_) / 2), vw_, vh_,
                                        poster_frame_image=str(poster) if poster.exists() else None, mime_type="video/mp4")
                    figs = []
                else:
                    text = text + [(f"[missing video {d['video']}: run «kit.py demo»]", False)]
            if d["anim"] and full_static:
                figs = []                       # the frame is already the slide background
            elif d["anim"]:
                auto = f"figures/static_{d['anim']}{suffix}.png"
                figs = [d["static"]] if d["static"] else ([auto] if (folder / auto).exists() else [])
                if not figs:
                    text = text + [("[Animation: see the HTML version. Run «kit.py statics» or add @static.]", False)]
            if kind == "number":
                box(sl, M, y0, 800 * px, height, [(d["number"] or "", False)], 84, c["accent"], font_title, True,
                    wrap=False, middle=True)
                box(sl, M + 860 * px, y0, width - 860 * px, height, text, 19, c["ink"], font_text, middle=True)
            elif figs and text:
                side = 560 * px
                for k, fig in enumerate(figs):
                    w = (width - side - 64 * px) / len(figs)
                    image(sl, fig, M + k * w, y0, w - 20 * px, height)
                box(sl, W - M - side, y0, side, height, text, 17, c["ink"], font_text)
            elif figs:
                w = width / len(figs)
                for k, fig in enumerate(figs):
                    image(sl, fig, M + k * w, y0, w - (40 * px if len(figs) > 1 else 0), height)
            else:   # text only: larger and vertically centered, like the HTML (deck.css)
                appendix = d["appendix"]
                box(sl, M, y0, width, height, text, 14 if appendix else 21, c["ink"], font_text, middle=not appendix)
        if d["appendix"]:
            box(sl, M, TOP - 40 * px, 400 * px, 40 * px, [("ANEXO" if es else "APPENDIX", False)], 10, c["ink_soft"], font_text)
        if d["source"]:
            box(sl, M, H - 95 * px, W - 2 * M - 120 * px, 60 * px, [(("Fuente: " if es else "Source: ") + d["source"], False)],
                11, c["ink_soft"], font_text)
        if profile.get("logo") and not template and kind != "cover":
            logo = profile["_dir"] / profile["logo"]
            if logo.exists() and logo.suffix.lower() in (".png", ".jpg", ".jpeg"):
                sl.shapes.add_picture(str(logo), int(W - 56 * px - 200 * px), int(40 * px), height=int(64 * px))
        if d["notes"]:
            sl.notes_slide.notes_text_frame.text = " ".join(_plain(n) for n in d["notes"])
    prs.core_properties.title = str(meta.get("title", ""))
    prs.core_properties.author = str(meta.get("author", ""))
    prs.save(out)
    print(f"[OK] {out}  ({len(slides)} slides; template layout: «{layout.name}»)")


# =====================================================================
# Commands
# =====================================================================

STARTER = {"en": """---
title: {title_yaml}
author: Your name
date: {date}
place:
profile: {profile}
lang: en
audience: who sees it, what they know and what they don't
message: the sentence they should remember in a week
---

<!-- One slide per «## Title». The title is the conclusion, as a sentence (PRINCIPLES.md, 4).
     Keys: @type cover|section|figure|two|number|text  @figure path  @source text  @number «10 cm/s»
     @anim name  @static path (final image of the animation, for PDF and .pptx)  @subtitle text
     @steps (reveals paragraphs and bullets one by one; «@steps 3» = levels of an animation)  @appendix
     @why (design note, key E)  @class full (full-bleed figure)  @kicker text (small line above the title)
     @video figures/demo.mp4 (demo video recorded with «kit.py demo»: plays when the slide opens)
     Front matter: «transition: fade» for a short fade; «motion: respect» to honor the OS reduce-motion setting.
     Lines with «- » are bullets; with «> » speaker notes; the rest are paragraphs. -->

## {title}
@type cover
@subtitle Subtitle

## Write the first conclusion here, as a sentence
@figure figures/example.png
@source where the evidence comes from
> What you say while this slide is on screen.
""", "es": """---
title: {title_yaml}
author: Tu nombre
date: {date}
place:
profile: {profile}
lang: es
audience: quién la ve, qué sabe y qué no
message: la frase que deben recordar en una semana
---

<!-- Una diapositiva por «## Título». El título es la conclusión, en una oración (PRINCIPLES.es.md, 4).
     Claves: @type cover|section|figure|two|number|text  @figure ruta  @source texto  @number «10 cm/s»
     @anim nombre  @static ruta  @subtitle texto  @steps  @appendix  @why  @class full  @kicker texto
     (también en español: @tipo portada|seccion|figura|dos|dato|texto, @figura, @fuente, @cifra, @pasos, @porque...)
     Líneas con «- » son viñetas; con «> » notas del orador; el resto, párrafos. -->

## {title}
@type cover
@subtitle Subtítulo

## Escribe aquí la primera conclusión, como oración
@figure figures/ejemplo.png
@source de dónde sale la evidencia
> Lo que dices mientras se ve esta diapositiva.
"""}


BRIEF = {"en": """# Brief: {title}

Answers to the intake questions (skill presentation-kit, "Start here"). Fill this in before writing the script.

## What the user wants to get (in this order)
- [ ] The talk itself first: what to say, slide by slide (speaker script)
- [ ] The content only: titles and evidence per slide (script.md), to review
- [ ] The finished deck
- [ ] A demo of software (see below)

## Format
- [ ] HTML (recommended: looks better, the only one with animations and demo videos)
- [ ] .pptx (someone else must edit it; animated slides become still images)

## Audience and message
- Audience, and what they already know:
- The one sentence they should remember in a week:
- Duration:
- Language:
- Style reference (deck, template, none):
- Register (sober / expressive):

## Demo of software (if any)
- What exactly to demonstrate (task, order, inputs):
- Real product or mock-up? (recorded from running code with «kit.py demo», or a labeled mock-up):
- How to start it (URL, command), data or test credentials:
""", "es": """# Brief: {title}

Respuestas a las preguntas iniciales (skill presentation-kit, «Start here»). Se completa antes de escribir el guion.

## Qué quiere obtener (en este orden)
- [ ] Primero la charla: qué decir, diapositiva por diapositiva (guion del orador)
- [ ] Solo el contenido: títulos y evidencia por diapositiva (script.es.md), para revisar
- [ ] La presentación terminada
- [ ] Una demostración de software (ver abajo)

## Formato
- [ ] HTML (recomendado: se ve mejor y es el único con animaciones y videos de demostración)
- [ ] .pptx (otra persona tiene que editarla; las diapositivas animadas quedan como imágenes fijas)

## Audiencia y mensaje
- Audiencia, y qué sabe ya:
- La frase que debe recordar en una semana:
- Duración:
- Idioma:
- Referencia de estilo (presentación, plantilla, ninguna):
- Registro (sobrio / expresivo):

## Demostración de software (si hay)
- Qué demostrar exactamente (tarea, orden, entradas):
- ¿Producto real o maqueta? (grabado desde el código que corre con «kit.py demo», o maqueta rotulada):
- Cómo se inicia (URL, comando), datos o credenciales de prueba:
"""}


def opt(args: list[str], name: str, default=None):
    return args[args.index(name) + 1] if name in args else default


def cmd_new(args: list[str]) -> None:
    import datetime
    import json
    if not args or args[0].startswith("-"):
        sys.exit('Usage: kit.py new <folder> [--profile sober] [--title "..."] [--lang es]')
    folder = Path(args[0]).resolve()
    profile = opt(args, "--profile", "sober")
    title = opt(args, "--title", folder.name.replace("_", " "))
    lang = opt(args, "--lang", "en")
    src = PROFILES / profile if (PROFILES / profile).is_dir() else Path(profile)
    if not (src / "profile.yml").exists():
        sys.exit(f"[ERROR] Profile «{profile}» not found. Profiles: {', '.join(p.name for p in PROFILES.iterdir() if p.is_dir())}")
    script = folder / ("script.md" if lang == "en" else f"script.{lang}.md")
    if script.exists():
        sys.exit(f"[ERROR] {script} already exists: not overwriting it.")
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "figures").mkdir(exist_ok=True)
    (folder / "engine").mkdir(exist_ok=True)
    for a in ("deck.css", "deck.js", "anim.js", "figures.py"):
        shutil.copy2(ENGINE / a, folder / "engine" / a)
    shutil.copytree(src, folder / "profile", dirs_exist_ok=True)
    script.write_text(STARTER["es" if lang == "es" else "en"].format(
        title=title, title_yaml=json.dumps(title, ensure_ascii=False), profile=src.name,
        date=datetime.date.today().isoformat()), encoding="utf-8")
    brief = folder / ("brief.md" if lang == "en" else f"brief.{lang}.md")
    if not brief.exists():
        brief.write_text(BRIEF["es" if lang == "es" else "en"].format(title=title), encoding="utf-8")
    print(f"[OK] new deck in {folder} with profile «{src.name}». Next: write {script.name} and run «kit.py build».")


def cmd_build(args: list[str]) -> None:
    if not args:
        sys.exit("Usage: kit.py build <folder> [--html] [--pptx] [--script script.es.md] [--overwrite]")
    folder = Path(args[0]).resolve()
    do_html = "--html" in args or "--pptx" not in args
    do_pptx = "--pptx" in args or "--html" not in args
    name = opt(args, "--script", "script.md")
    if name == "script.md" and not (folder / name).exists() and (folder / "guion.md").exists():
        name = "guion.md"
    suf = script_suffix(name)
    meta, slides = read_script(folder / name)
    profile = read_profile(folder / "profile")
    if do_html:
        dest = folder / f"deck{suf}.html"
        if dest.exists() and "Generated from the script" not in dest.read_text(encoding="utf-8")[:300] \
                and "--overwrite" not in args:
            dest = folder / f"draft{suf}.html"
            print("[WARNING] the deck exists and was not generated by the kit (may be hand-made): writing a draft.")
        dest.write_text(html_from_script(meta, slides, profile.get("logo"), folder), encoding="utf-8")
        print(f"[OK] {dest}  ({len(slides)} slides)")
    if do_pptx:
        pptx_from_script(meta, slides, profile, folder, folder / f"{folder.name}{suf}.pptx", suf)


def _data_uri(path: Path) -> str:
    kind = mimetypes.guess_type(path.name)[0] or {".woff2": "font/woff2", ".woff": "font/woff"}.get(path.suffix, "application/octet-stream")
    return f"data:{kind};base64,{base64.b64encode(path.read_bytes()).decode()}"


def cmd_package(args: list[str]) -> None:
    if not args:
        sys.exit("Usage: kit.py package <deck.html> [out.html]")
    src = Path(args[0]).resolve()
    base = src.parent
    # deck.html -> <folder>.html ; deck.es.html -> <folder>.es.html
    out = Path(args[1]).resolve() if len(args) > 1 else base / f"{base.name}{src.name[4:-5]}.html"
    text = src.read_text(encoding="utf-8")
    missing = []

    def css_inline(m):
        path = (base / m.group(1)).resolve()
        if not path.exists():
            missing.append(m.group(1)); return m.group(0)
        css = path.read_text(encoding="utf-8")
        def url(mu):
            u = mu.group(1).strip("'\"")
            if u.startswith(("data:", "http:", "https:", "#")):
                return mu.group(0)
            r = (path.parent / u).resolve()
            if not r.exists():
                missing.append(u); return mu.group(0)
            return f"url('{_data_uri(r)}')"
        return "<style>\n" + re.sub(r"url\(([^)]+)\)", url, css) + "\n</style>"

    def js_inline(m):
        path = (base / m.group(1)).resolve()
        if not path.exists():
            missing.append(m.group(1)); return m.group(0)
        return "<script>\n" + path.read_text(encoding="utf-8").replace("</script", "<\\/script") + "\n</script>"

    def attribute(m):
        attr, u = m.group(1), m.group(2)
        if u.startswith(("data:", "http:", "https:", "#")) or not u:
            return m.group(0)
        r = (base / u).resolve()
        if not r.exists():
            missing.append(u); return m.group(0)
        return f'{attr}="{_data_uri(r)}"'

    text = re.sub(r'<link[^>]+rel="stylesheet"[^>]+href="([^"]+)"[^>]*>', css_inline, text)
    text = re.sub(r'<script src="([^"]+)"></script>', js_inline, text)
    text = re.sub(r'\b(src|data-logo|href|poster)="([^"]+\.(?:png|jpe?g|svg|gif|webp|mp4|webm))"', attribute, text)
    out.write_text(text, encoding="utf-8")
    print(f"[OK] {out}  ({out.stat().st_size / 1e6:.1f} MB, single file)")
    for f in missing:
        print(f"[WARNING] {f} not found: left as an external reference.")


def _count_slides(path: Path) -> int:
    """Slides in the HTML, ignoring examples that appear inside comments of packaged CSS or JS."""
    text = path.read_text(encoding="utf-8")
    text = re.sub(r"<script.*?</script>|<style.*?</style>|<!--.*?-->", "", text, flags=re.S)
    return len(re.findall(r'<section class="slide[ "]', text))


def cmd_pdf(args: list[str]) -> None:
    pos = [a for a in args if not a.startswith("--")]
    if not pos:
        sys.exit("Usage: kit.py pdf <file.html> [out.pdf] [--overwrite]")
    src = Path(pos[0]).resolve()
    pdf = Path(pos[1]).resolve() if len(pos) > 1 else src.with_suffix(".pdf")
    if pdf.exists() and "--overwrite" not in args:
        sys.exit(f"[ERROR] {pdf.name} already exists. Use --overwrite to replace it.")
    n_slides = _count_slides(src)
    profile_dir = Path(tempfile.gettempdir()) / "presentation-kit-browser"
    if pdf.exists():
        pdf.unlink()
    r = subprocess.run([browser(), "--headless=new", "--disable-gpu", "--no-pdf-header-footer", "--hide-scrollbars",
                        f"--user-data-dir={profile_dir}", "--virtual-time-budget=8000", f"--print-to-pdf={pdf}",
                        src.as_uri() + "?print"], capture_output=True, text=True, timeout=240)
    if not pdf.exists() or pdf.stat().st_size == 0:
        sys.exit(f"[ERROR] The browser did not produce the PDF: {(r.stderr or r.stdout)[-400:]}")
    pages = len(re.findall(rb"/Type\s*/Page(?!s)", pdf.read_bytes()))
    print(f"[OK] {pdf}  ({pages} pages; the deck has {n_slides} slides)")
    if pages != n_slides:
        print("[WARNING] Page count and slide count differ: open the PDF and check it.")


LABELS = {"introduction", "background", "methods", "methodology", "results", "conclusions", "conclusion", "agenda",
          "outline", "objectives", "thanks", "thank you", "questions", "summary", "motivation", "discussion",
          "future work", "next steps", "introducción", "introduccion", "contexto", "metodología", "metodologia",
          "resultados", "conclusiones", "conclusión", "índice", "indice", "objetivos", "gracias", "¡gracias!",
          "preguntas", "resumen", "motivación", "motivacion", "antecedentes", "discusión", "discusion",
          "trabajo futuro", "próximos pasos"}


def cmd_review(args: list[str]) -> None:
    if not args:
        sys.exit("Usage: kit.py review <deck.html>")
    path = Path(args[0]).resolve()
    text = path.read_text(encoding="utf-8")
    sections = re.findall(r'(<section class="slide[^"]*"[^>]*>)(.*?)</section>', text, re.S)
    code = text + "".join((path.parent / s).read_text(encoding="utf-8")
                          for s in re.findall(r'<script src="([^"]+)"', text)
                          if not s.startswith("http") and (path.parent / s).exists())
    findings, titles = [], []
    for i, (opening, body) in enumerate(sections, 1):
        cls = re.search(r'class="([^"]+)"', opening).group(1)
        h1 = re.search(r"<h1[^>]*>(.*?)</h1>", body, re.S)
        t = re.sub(r"<[^>]+>", "", h1.group(1)).strip() if h1 else ""
        titles.append(t)
        special = any(k in cls.split() for k in ("cover", "section", "appendix"))
        if not t:
            findings.append((i, "has no title (h1)"))
        elif not special and (t.lower().strip(".:") in LABELS or len(t.split()) < 3):
            findings.append((i, f"the title «{t}» is a label, not a conclusion (principle 4)"))
        # On screen, notes, «why it works» and the source line do not count (small print apart).
        visible = re.sub(r'<aside class="(notes|why)">.*?</aside>|<p class="source[^"]*">.*?</p>', "", body, flags=re.S)
        visible = re.sub(r"<script.*?</script>|<style.*?</style>|<[^>]+>", " ", visible, flags=re.S)
        words = len(visible.split())
        if not special and words > 75:
            findings.append((i, f"has {words} words on screen: is it a slide or a document? (principle 5)"))
        if re.search(r"<(img|svg|figure|canvas)\b", body) and 'class="source' not in body and not special:
            findings.append((i, "has a figure and no source line (principle 8)"))
        for src in re.findall(r'<img[^>]+src="([^"]+)"', body):
            if not src.startswith(("data:", "http")) and not (path.parent / src).exists():
                findings.append((i, f"image {src} does not exist"))
        anim = re.search(r'data-anim="([^"]+)"', opening)
        if anim and not re.search(rf"\.{re.escape(anim.group(1))}\s*=|\b{re.escape(anim.group(1))}\s*:", code):
            findings.append((i, f"asks for animation «{anim.group(1)}» and its function was not found"))
        if '<aside class="notes">' not in body and not special:
            findings.append((i, "has no speaker notes (not an error; useful for rehearsing)"))
    print(f"Mechanical review of {path.name}: {len(sections)} slides\n")
    print("The titles, read in a row (do they tell the story on their own?):")
    for i, t in enumerate(titles, 1):
        print(f"  {i:>2}. {t}")
    print()
    if not findings:
        print("[OK] No mechanical findings. Still to do: review by eye (PRINCIPLES.md, «Review checklist»).")
    for i, f in findings:
        print(f"  slide {i}: {f}")
    print("\nThis checks what can be counted. The rest of the checklist is checked by looking.")


def cmd_profile(args: list[str]) -> None:
    if "--regenerate" in args:
        regenerate_profile(Path(opt(args, "--regenerate")).resolve())
        return
    if "--from" not in args or "--name" not in args:
        sys.exit("Usage: kit.py profile --from <ref.pptx> --name <name> [--dest profiles]  |  "
                 "kit.py profile --regenerate <profile_dir>")
    ref = Path(opt(args, "--from")).resolve()
    dest = Path(opt(args, "--dest")).resolve() if "--dest" in args else PROFILES
    if ref.suffix.lower() not in (".pptx", ".potx"):
        sys.exit("[ERROR] Only .pptx or .potx can be read. For a PDF, an image or a web page: copy profiles/sober, edit "
                 "profile.yml while looking at the reference and run «kit.py profile --regenerate».")
    profile_from_pptx(ref, opt(args, "--name"), dest)


def cmd_shots(args: list[str]) -> None:
    """One image per slide of the HTML, headless: to review by eye (or to let a model review)."""
    if len(args) < 2:
        sys.exit("Usage: kit.py shots <deck.html> <png_dir> [--only 3,5,7]")
    src, folder = Path(args[0]).resolve(), Path(args[1]).resolve()
    folder.mkdir(parents=True, exist_ok=True)
    exe = browser()
    n = _count_slides(src)
    only = [int(x) for x in opt(args, "--only").split(",")] if "--only" in args else range(1, n + 1)
    profile_dir = Path(tempfile.gettempdir()) / "presentation-kit-browser"
    for k in only:
        dest = folder / f"slide_{k:02d}.png"
        subprocess.run([exe, "--headless=new", "--disable-gpu", "--hide-scrollbars", f"--user-data-dir={profile_dir}",
                        "--window-size=1920,1080", "--virtual-time-budget=4000", f"--screenshot={dest}",
                        src.as_uri() + f"?print#{k}"], capture_output=True, timeout=120)
        if not dest.exists():
            print(f"[WARNING] no screenshot for slide {k}")
    print(f"[OK] {len(list(only))} screenshots in {folder} (each slide in its final state)")


def cmd_statics(args: list[str]) -> None:
    """Photographs the final state of each animation (without title or source) as figures/static_<name>.png.
    The .pptx uses them instead of the animation; the PDF does not need them (it prints the final SVG state)."""
    if not args:
        sys.exit("Usage: kit.py statics <folder> [--script script.es.md]")
    folder = Path(args[0]).resolve()
    suf = script_suffix(opt(args, "--script")) if "--script" in args else ""
    src = folder / f"deck{suf}.html"
    text = re.sub(r"<script.*?</script>|<!--.*?-->", "", src.read_text(encoding="utf-8"), flags=re.S)
    exe = browser()
    profile_dir = Path(tempfile.gettempdir()) / "presentation-kit-browser"
    done = 0
    for k, attrs in enumerate(re.findall(r'<section class="slide[^"]*"([^>]*)>', text), 1):
        m = re.search(r'data-anim="([^"]+)"', attrs)
        if not m:
            continue
        dest = folder / "figures" / f"static_{m.group(1)}{suf}.png"
        dest.parent.mkdir(exist_ok=True)
        subprocess.run([exe, "--headless=new", "--disable-gpu", "--hide-scrollbars", f"--user-data-dir={profile_dir}",
                        "--window-size=1920,1080", "--virtual-time-budget=5000", f"--screenshot={dest}",
                        src.as_uri() + f"?print&scene-only#{k}"], capture_output=True, timeout=120)
        done += dest.exists()
        print(f"  slide {k}: {dest.name}" if dest.exists() else f"[WARNING] slide {k} failed")
    print(f"[OK] {done} statics in {folder / 'figures'}. Run «build --pptx» again to use them.")


def cmd_demo(args: list[str]) -> None:
    """Records a demo video from the real running interface (engine/demo.py)."""
    if not args:
        sys.exit("Usage: kit.py demo <demo.yml>   (see engine/demo.py for the format)")
    sys.path.insert(0, str(ENGINE))
    import demo
    demo.main(args)


GUIDE_YML = """# Study guide: one self-contained HTML file per page, opened with a double click (no server, no internet).
title: {title_yaml}
subtitle: What the reader will be able to do after the guide
answer: The answer first, in one or two sentences; the pages explain it.
lang: {lang}
profile: {profile}   # only the accent colors are taken from it
anim: false          # true also loads engine/anim.js (K.scene) for animated SVG figures
# Pages: the files in pages/, in name order. Each starts with <!--META {{"title": "...", "what": "..."}} META-->.
# A data.json next to this file reaches every page as G.data.
"""

GUIDE_PAGE = """<!--META {{"title": "{title}", "what": "The question this page answers"}} META-->
<p class="kicker">Part 1</p>
<h1>A title that states what the page shows</h1>
<p class="answer">The answer, before the detail.</p>

<div class="predict" data-correct="b">
  <p class="q">Before you look: what do you expect?</p>
  <div class="options"><button data-option="a">Option A</button><button data-option="b">Option B</button></div>
  <p class="why">Why B: one or two sentences.</p>
</div>
<div class="sim wide" data-gate>
  <div class="controls"><label>N <input type="range" id="n" min="1" max="100" value="10" data-format="int"><output for="n"></output></label></div>
  <svg id="plot"></svg>
</div>
<script>
G.bind("n", n => G.plot(document.getElementById("plot"), {{
  xlabel: "x (unit)", ylabel: "y (unit)",
  series: [{{ x: [0, 1, 2], y: [0, n, 2 * n] }}],
}}));
</script>

<details class="check"><summary>Check yourself: a question</summary><p>Its answer.</p></details>
"""


def _guide_meta(text: str) -> tuple[dict, str]:
    m = re.match(r"\s*<!--META\s*(\{.*?\})\s*META-->", text, re.S)
    if not m:
        return {}, text
    import json
    return json.loads(m.group(1)), text[m.end():]


def cmd_guide(args: list[str]) -> None:
    """Study guide: pages/*.html fragments → one self-contained HTML per page plus index.html (the map)."""
    if not args:
        sys.exit("Usage: kit.py guide new <folder> [--title t] [--lang es]   |   kit.py guide <folder>")
    if args[0] == "new":
        if len(args) < 2 or args[1].startswith("-"):
            sys.exit("Usage: kit.py guide new <folder> [--title t] [--lang es]")
        folder = Path(args[1]).resolve()
        if folder.exists() and any(folder.iterdir()):
            sys.exit(f"[ERROR] {folder} already exists and is not empty.")
        import json
        title = opt(args, "--title", "Study guide")
        lang = opt(args, "--lang", "en")
        (folder / "pages").mkdir(parents=True)
        (folder / "guide.yml").write_text(GUIDE_YML.format(title_yaml=json.dumps(title, ensure_ascii=False), lang=lang, profile=(KIT / "profiles" / "sober").as_posix()), encoding="utf-8")
        (folder / "pages" / "01_first.html").write_text(GUIDE_PAGE.format(title="First page"), encoding="utf-8")
        print(f"[OK] guide in {folder}: edit guide.yml and pages/, then «kit.py guide {folder}».")
        return
    import json
    folder = Path(args[0]).resolve()
    spec = yaml.safe_load((folder / "guide.yml").read_text(encoding="utf-8")) or {}
    lang = spec.get("lang", "en")
    t = (lambda en, es: es if lang == "es" else en)
    pages = sorted((folder / "pages").glob("*.html"))
    if not pages:
        sys.exit(f"[ERROR] no pages in {folder / 'pages'}")
    css = (ENGINE / "guide.css").read_text(encoding="utf-8")
    if spec.get("profile"):
        prof = Path(spec["profile"])
        prof = prof if prof.is_absolute() else (folder / prof)
        c = yaml.safe_load((prof / "profile.yml").read_text(encoding="utf-8"))["colors"]
        a, a2 = c["accent"], c["accent_2"]
        dark = f"--accent: color-mix(in srgb, {a} 55%, white); --accent-2: color-mix(in srgb, {a2} 70%, white);"
        css += (f"\n:root {{ --accent: {a}; --accent-2: {a2}; }}\n"
                f"@media (prefers-color-scheme: dark) {{ :root:not([data-theme=\"light\"]) {{ {dark} }} }}\n"
                f":root[data-theme=\"dark\"] {{ {dark} }}\n")
    js = ""
    if (folder / "data.json").exists():
        js += "window.GUIDE_DATA = " + (folder / "data.json").read_text(encoding="utf-8") + ";\n"
    if spec.get("anim"):
        js += (ENGINE / "anim.js").read_text(encoding="utf-8") + "\n"
    js += (ENGINE / "guide.js").read_text(encoding="utf-8")
    js = js.replace("</script", "<\\/script")
    metas = []
    for p in pages:
        meta, body = _guide_meta(p.read_text(encoding="utf-8"))
        metas.append((p.stem, meta.get("title", p.stem), meta.get("what", ""), body))

    def inline(body: str, missing: list) -> str:
        def attribute(m):
            u = m.group(2)
            if u.startswith(("data:", "http:", "https:", "#")):
                return m.group(0)
            r = (folder / u).resolve()
            if not r.exists():
                missing.append(u); return m.group(0)
            return f'{m.group(1)}="{_data_uri(r)}"'
        return re.sub(r'\b(src|href|poster)="([^"]+\.(?:png|jpe?g|svg|gif|webp|mp4|webm))"', attribute, body)

    def page(title: str, main: str) -> str:
        bar = (f'<nav class="bar"><a class="guide-title" href="index.html">{html.escape(spec.get("title", ""))}</a>'
               f'<span class="spacer"></span>'
               f'<button id="exercise-btn" aria-pressed="false" title="{t("Lock each simulator until you predict", "Bloquea cada simulador hasta que predigas")}">'
               f'{t("Exercise mode", "Modo ejercicio")}</button>'
               f'<button id="theme-btn">{t("Light / dark", "Claro / oscuro")}</button></nav>')
        return (f'<!doctype html>\n<html lang="{lang}">\n<head>\n<meta charset="utf-8">\n'
                f'<meta name="viewport" content="width=device-width, initial-scale=1">\n<title>{html.escape(title)}</title>\n'
                f'<style>\n{css}\n</style>\n<script>\n{js}\n</script>\n</head>\n<body>\n{bar}\n<main>\n{main}\n</main>\n</body>\n</html>\n')

    missing: list[str] = []
    for i, (stem, title, _, body) in enumerate(metas):
        prev = f'<a href="{metas[i - 1][0]}.html">← {html.escape(metas[i - 1][1])}</a>' if i else f'<a href="index.html">← {t("Map", "Mapa")}</a>'
        nxt = f'<a href="{metas[i + 1][0]}.html">{html.escape(metas[i + 1][1])} →</a>' if i + 1 < len(metas) else "<span></span>"
        main = inline(body, missing) + f'\n<nav class="pager">{prev}{nxt}</nav>'
        (folder / f"{stem}.html").write_text(page(f'{title} · {spec.get("title", "")}', main), encoding="utf-8")
    items = "\n".join(f'<li><a href="{s}.html">{i}. {html.escape(ti)}</a><span class="what">{html.escape(w)}</span></li>'
                      for i, (s, ti, w, _) in enumerate(metas, 1))
    index = (f'<h1>{html.escape(spec.get("title", ""))}</h1>\n'
             + (f'<p class="lead">{html.escape(spec["subtitle"])}</p>\n' if spec.get("subtitle") else "")
             + (f'<p class="answer">{html.escape(spec["answer"])}</p>\n' if spec.get("answer") else "")
             + f'<ol class="map">\n{items}\n</ol>\n'
             + (f'<p class="limits">{html.escape(spec["limits"])}</p>\n' if spec.get("limits") else ""))
    (folder / "index.html").write_text(page(spec.get("title", ""), inline(index, missing)), encoding="utf-8")
    print(f"[OK] {len(metas)} pages + index.html in {folder} (each opens with a double click)")
    for f in sorted(set(missing)):
        print(f"[WARNING] {f} not found: left as an external reference.")


def cmd_render(args: list[str]) -> None:
    if len(args) < 2:
        sys.exit("Usage: kit.py render <file.pptx> <png_dir>")
    src, folder = Path(args[0]).resolve(), Path(args[1]).resolve()
    folder.mkdir(parents=True, exist_ok=True)
    ps = f"""
$pp = New-Object -ComObject PowerPoint.Application
$pres = $pp.Presentations.Open('{src}', $true, $false, $false)
$i = 1
foreach ($s in $pres.Slides) {{ $s.Export((Join-Path '{folder}' ('slide_{{0:D2}}.png' -f $i)), 'PNG', 1600, 900); $i++ }}
$pres.Close(); $pp.Quit()
"""
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)
    print(f"[OK] images in {folder}")


COMMANDS = {"new": cmd_new, "build": cmd_build, "package": cmd_package, "pdf": cmd_pdf, "review": cmd_review,
            "profile": cmd_profile, "render": cmd_render, "shots": cmd_shots, "statics": cmd_statics, "demo": cmd_demo,
            "guide": cmd_guide}
ALIASES_ES = {"nueva": "new", "construir": "build", "empaquetar": "package", "revisar": "review", "perfil": "profile",
              "renderizar": "render", "capturas": "shots", "estaticos": "statics", "guia": "guide"}
OPTIONS_ES = {"--perfil": "--profile", "--titulo": "--title", "--sobrescribir": "--overwrite", "--solo": "--only",
              "--desde": "--from", "--nombre": "--name", "--destino": "--dest", "--regenerar": "--regenerate",
              "--guion": "--script", "--idioma": "--lang"}

if __name__ == "__main__":
    command = ALIASES_ES.get(sys.argv[1], sys.argv[1]) if len(sys.argv) > 1 else ""
    if command not in COMMANDS:
        print(__doc__)
        sys.exit(2)
    COMMANDS[command]([OPTIONS_ES.get(a, a) for a in sys.argv[2:]])
